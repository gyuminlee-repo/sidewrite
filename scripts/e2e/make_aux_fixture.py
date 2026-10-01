#!/usr/bin/env python3
"""Fake manuscript for the stage 4 acceptance run (Tools tab).

Every sentence, name and number is invented. Two variants:

  current   tests/fixtures/aux_manuscript.docx     the manuscript the pane opens
  v1        tests/fixtures/aux_manuscript_v1.docx  an earlier version, the
                                                   baseline of the marked copy

Planted in both, one each, for the format check (profile PROFILE below):
  a statistics symbol "p" that is not italic
  a species name that is not italic
  a formula whose digit is not subscript
  one body paragraph without the first-line indent of the others
  an empty table
and for the count check: an abstract of more than 250 words under the
heading "Abstract", and a Keywords line with three keywords.

The earlier version differs from the current one in two sentences and lacks
the last body paragraph.

Usage: python3 scripts/e2e/make_aux_fixture.py [--variant v1] [output.docx]
       python3 scripts/e2e/make_aux_fixture.py --profile   (prints PROFILE as JSON)
       python3 scripts/e2e/make_aux_fixture.py --dump x.docx   (paragraphs as JSON, for tests)
"""
import datetime
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

FIXTURES = Path(__file__).resolve().parents[2] / "tests" / "fixtures"
OUT = {"current": FIXTURES / "aux_manuscript.docx", "v1": FIXTURES / "aux_manuscript_v1.docx"}
FIXED_TIME = datetime.datetime(2026, 1, 1, 0, 0, 0)
INDENT = Pt(36)

GENUS, EPITHET, FORMULA = "Examplea", "fictiva", "CO2"
SPECIES = "%s %s" % (GENUS, EPITHET)

# The journal profile the acceptance run puts next to the manuscript: the
# format rules that find the planted violations.
PROFILE = {
    "name": "s4-acceptance",
    "italics": {"statistics": True, "genera": [GENUS], "epithets": [EPITHET]},
    "subscripts": [FORMULA],
    "first_line_indent": "consistent",
}
# What the format check must report: rule -> number of findings.
PLANTED = {"stat": 1, "species": 1, "subscript": 1, "indent": 1, "empty_table": 1}
ABSTRACT_LIMIT = 250

TITLE = "Growth of an invented bacterium in a shaken flask"
ABSTRACT_SENTENCE = (
    "The invented cultures of this placeholder study were grown in shaken flasks and sampled "
    "at regular intervals so that a software test has an abstract of a known length to count. "
)
KEYWORDS = "Keywords: flask, growth, placeholder"

# Body paragraphs: a list of (text, format) pieces; format is "" (plain),
# "i" (italic) or "sub" (subscript). Every paragraph has more than twenty
# words and ends a sentence, so the indent rule judges it.
CLEAN_1 = [
    ("The model organism of this invented study, ", ""),
    (SPECIES, "i"),
    (", was kept on agar plates and moved to liquid medium two days before each run, and the "
     "flasks were flushed with air that held a fixed share of CO", ""),
    ("2", "sub"),
    (" before they were sealed.", ""),
]
STAT = [
    ("Growth in the two invented conditions was compared at the end of the run, and the "
     "difference between the final optical densities was significant (p < 0.05) in every one "
     "of the three repeats.", ""),
]
SPECIES_PLAIN = [
    ("Cells of %s taken from the plateau were washed twice in buffer and kept on ice, and "
     "their protein content was estimated with a dye-binding assay against a standard "
     "curve." % SPECIES, ""),
]
FORMULA_PLAIN = [
    ("The headspace of every flask was sampled once per day with a gas-tight syringe, and the "
     "share of %s in it was read on a detector that had been calibrated on the same "
     "morning." % FORMULA, ""),
]
FLUSH = [
    ("Samples for the dry weight were drawn at the same time as the optical density readings, "
     "filtered, dried overnight and weighed on a balance whose drift was checked before every "
     "series.", ""),
]
CLEAN_2 = [
    ("Each run was repeated three times on different days with fresh medium, and the mean and "
     "the standard deviation of the three repeats are what the invented figures of this "
     "placeholder text would show.", ""),
]
CLEAN_3 = [
    ("The spread between the repeats stayed small during exponential growth and widened "
     "slightly at the plateau, which does not change the conclusion of this illustrative "
     "example in any way.", ""),
]
LAST = [
    ("A final invented paragraph was added in the later version of this text, so that a "
     "comparison of the two versions has one paragraph that exists only in the newer file "
     "of the pair.", ""),
]
# Wording of the earlier version.
V1_REWRITES = {
    "filtered, dried overnight and weighed": "filtered, dried for two hours and weighed",
    "repeated three times on different days": "repeated twice on the same day",
}

CAPTION = "Table 1. Invented plate layout of the placeholder run."
CLOSING = "No real data are in this document."


def add_body(doc, pieces, indent, variant):
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = indent
    for text, fmt in pieces:
        if variant == "v1":
            for old, new in V1_REWRITES.items():
                text = text.replace(old, new)
        run = para.add_run(text)
        if fmt == "i":
            run.italic = True
        elif fmt == "sub":
            run.font.subscript = True
    return para


def abstract_text():
    words = (ABSTRACT_SENTENCE * 12).split()
    return " ".join(words[: ABSTRACT_LIMIT + 20]) + "."


def build(out_path, variant="current"):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    cp = doc.core_properties
    cp.author = "Fixture Author"
    cp.last_modified_by = "Fixture Author"
    cp.title = TITLE
    cp.created = FIXED_TIME
    cp.modified = FIXED_TIME
    cp.revision = 1

    doc.add_paragraph(TITLE, style="Title")
    doc.add_paragraph("Abstract", style="Heading 1")
    doc.add_paragraph(abstract_text()).paragraph_format.first_line_indent = INDENT
    doc.add_paragraph(KEYWORDS)
    doc.add_paragraph("Methods", style="Heading 1")
    for pieces in (CLEAN_1, STAT, SPECIES_PLAIN, FORMULA_PLAIN):
        add_body(doc, pieces, INDENT, variant)
    add_body(doc, FLUSH, Pt(0), variant)
    doc.add_paragraph("Results", style="Heading 1")
    for pieces in (CLEAN_2, CLEAN_3):
        add_body(doc, pieces, INDENT, variant)
    if variant != "v1":
        add_body(doc, LAST, INDENT, variant)
    doc.add_paragraph(CAPTION)
    doc.add_table(rows=2, cols=2)
    doc.add_paragraph(CLOSING)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def dump(path):
    """The paragraphs of a docx as the format check needs them, for tests:
    [{ xml, text, first_line, table_level, heading }] in document order."""
    from docx.oxml.ns import qn
    from lxml import etree

    out = []
    for p in Document(path).element.body.iter(qn("w:p")):
        ind = p.find("%s/%s" % (qn("w:pPr"), qn("w:ind")))
        first = ind.get(qn("w:firstLine")) if ind is not None else None
        style = p.find("%s/%s" % (qn("w:pPr"), qn("w:pStyle")))
        out.append({
            "xml": etree.tostring(p, encoding="unicode"),
            "text": "".join(t.text or "" for t in p.iter(qn("w:t"))),
            "first_line": int(first) / 20 if first else 0,
            "table_level": sum(1 for a in p.iterancestors(qn("w:tbl"))),
            "heading": style is not None and style.get(qn("w:val")).startswith("Heading"),
        })
    return out


def main(argv):
    args = argv[1:]
    if args[:1] == ["--profile"]:
        print(json.dumps(PROFILE))
        return
    if args[:1] == ["--dump"]:
        print(json.dumps(dump(args[1])))
        return
    variant = "current"
    if args[:1] == ["--variant"]:
        variant, args = args[1], args[2:]
    if variant not in OUT:
        sys.exit("unknown variant %s (current, v1)" % variant)
    print(build(Path(args[0]) if args else OUT[variant], variant))


if __name__ == "__main__":
    main(sys.argv)
