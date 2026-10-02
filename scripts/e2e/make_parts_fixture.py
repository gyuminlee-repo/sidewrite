#!/usr/bin/env python3
"""Fake manuscript and fake supplementary file for the `parts` acceptance stage.

Every sentence, name and number is invented. Two documents that belong to one
revision:

  paper  tests/fixtures/parts_paper.docx  a short manuscript (in the run folder: paper.docx)
  si     tests/fixtures/parts_si.docx     supplementary methods and one table (SI.docx)

and the review that goes with them, tests/fixtures/parts_review.txt, whose four
comments ask for a change

  R1.1  only in the supplementary methods (the incubation temperature)
  R1.2  in both files (the shaking speed)
  R1.3  only in the manuscript (a word in the Results)
  R1.4  in a figure file (the unit on an axis of Figure 1)

What the stage relies on (tests/test_make_parts_fixture.py checks it):
  - the sentence each comment is about is in the file it names, once;
  - the word "novel" is in neither document (the ledger of the stage,
    tests/fixtures/parts_ledger.json, forbids it);
  - GUARD["find"] is in the supplementary file exactly once: the stage sends an
    edit that turns it into GUARD["replace"], which holds the forbidden word.

Usage: python3 scripts/e2e/make_parts_fixture.py            (writes both fixtures)
       python3 scripts/e2e/make_parts_fixture.py paper out.docx
       python3 scripts/e2e/make_parts_fixture.py si out.docx
       python3 scripts/e2e/make_parts_fixture.py --spec     (what the stage reads, as JSON)
       python3 scripts/e2e/make_parts_fixture.py --dump x.docx   (paragraph texts as JSON)

Writing a fixture needs python-docx (a development tool; the fixtures are
committed). --spec and --dump use the standard library only.
"""
import datetime
import json
import sys
import zipfile
from pathlib import Path

FIXTURES = Path(__file__).resolve().parents[2] / "tests" / "fixtures"
OUT = {"paper": FIXTURES / "parts_paper.docx", "si": FIXTURES / "parts_si.docx"}
FIXED_TIME = datetime.datetime(2026, 1, 1, 0, 0, 0)

# (style, text). style: "title", "heading" or "body".
PAPER = [
    ("title", "Growth of an invented bacterium on an invented carbon source"),
    ("heading", "Abstract"),
    (
        "body",
        "An invented bacterium was grown in shaken flasks on an invented carbon source. "
        "This placeholder text exists so that a software test has a manuscript to edit.",
    ),
    ("heading", "Methods"),
    ("body", "Starter cultures were grown in an invented minimal medium."),
    ("body", "The main cultures were shaken at a constant speed."),
    ("body", "Optical density was read once every hour."),
    ("heading", "Results"),
    (
        "body",
        "The invented strain grew good on the invented carbon source and reached a plateau "
        "after twelve hours.",
    ),
    ("body", "Figure 1 shows the growth curve of the invented strain."),
    ("heading", "Supplementary material"),
    ("body", "The Supplementary Methods and Supplementary Table S1 are in a separate file."),
]
SI = [
    ("title", "Supplementary Information"),
    ("heading", "Supplementary Methods"),
    ("body", "Starter cultures were incubated overnight before they were used."),
    ("body", "The starter flasks were shaken at a constant speed."),
    ("body", "The invented carbon source was added to each flask at the start."),
    ("heading", "Supplementary Table S1"),
    ("body", "Conditions of the invented starter cultures."),
]
SI_TABLE = [
    ("Condition", "Value"),
    ("Incubation temperature", "30 degrees Celsius"),
    ("Culture volume", "20 mL"),
]

# The edit the stage sends to the supplementary file to show the forbid guard.
GUARD = {"ref": "R1.1", "find": "invented carbon source", "replace": "novel carbon source"}

# What the stage reads: the file names in the run folder, which comment is of
# which kind, the requests the stage types for each comment (so their ids do
# not depend on a model), and the guard edit.
SPEC = {
    "names": {"paper": "paper.docx", "si": "SI.docx"},
    "refs": {"si_only": "R1.1", "both": "R1.2", "main_only": "R1.3", "figure": "R1.4"},
    "requirements": {
        "R1.1": [
            "state in the Supplementary Methods that the starter cultures were incubated at 30 degrees Celsius"
        ],
        "R1.2": [
            "give the shaking speed of 180 rpm in the Methods of the manuscript",
            "give the shaking speed of 180 rpm in the Supplementary Methods",
        ],
        "R1.3": ["write grew well instead of grew good in the Results"],
        "R1.4": ["add the unit to the vertical axis of Figure 1"],
    },
    "guard": GUARD,
    "forbidden_word": "novel",
}


def build(variant, out_path=None):
    from docx import Document

    rows = PAPER if variant == "paper" else SI
    doc = Document()
    for style, text in rows:
        if style == "title":
            doc.add_heading(text, level=0)
        elif style == "heading":
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    if variant == "si":
        table = doc.add_table(rows=len(SI_TABLE), cols=2)
        for r, cells in enumerate(SI_TABLE):
            for c, text in enumerate(cells):
                table.cell(r, c).text = text
    props = doc.core_properties
    props.author = "Fixture"
    props.last_modified_by = "Fixture"
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    out_path = Path(out_path or OUT[variant])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def dump(path):
    """The text of every paragraph of the document, table cells included, in
    document order. Standard library only."""
    from xml.etree import ElementTree as ET

    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    return ["".join(t.text or "" for t in p.iter(ns + "t")) for p in root.iter(ns + "p")]


if __name__ == "__main__":
    args = sys.argv[1:]
    if args[:1] == ["--spec"]:
        print(json.dumps(SPEC))
    elif args[:1] == ["--dump"] and len(args) == 2:
        print(json.dumps(dump(args[1])))
    elif not args:
        for name in ("paper", "si"):
            print(build(name))
    elif args[0] in OUT and len(args) <= 2:
        print(build(*args))
    else:
        sys.exit(__doc__)
