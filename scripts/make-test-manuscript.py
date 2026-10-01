#!/usr/bin/env python3
"""Generate a public, fake scientific manuscript for the stage-2 spike.

The document has three Heading 1 sections and Word line numbering turned on
(w:lnNumType countBy=1 restart=continuous) so a Word-rendered PDF carries
margin line numbers for scripts/pdf-line-map.py.

All text is invented filler. It names no real study, person, or institution.

Usage: python3 scripts/make-test-manuscript.py [--variant review] [--compat 14] [output.docx]
Default output: tests/fixtures/spike_manuscript.docx
(--variant review: tests/fixtures/review_manuscript.docx)

The document is written in compatibility mode 15 (Word 2013 and later). The
python-docx default template stores mode 14 (Word 2010), and in that mode
Word refuses to mark a comment resolved. --compat 14 keeps the old mode for
a test of that refusal; give it an output path outside tests/fixtures.

The review variant is the same text with ledger violations planted for the
stage-2 acceptance test (see PLANTS) and one comment by the author.
"""

import datetime
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

FIXTURES = Path(__file__).resolve().parent.parent / "tests" / "fixtures"
DEFAULT_OUT = FIXTURES / "spike_manuscript.docx"
REVIEW_OUT = FIXTURES / "review_manuscript.docx"

# Review variant: (section number in SECTIONS, paragraph position, sentence
# appended to the first copy). Planted: "in-vitro" twice (ledger L1 violations), "in
# vitro" once written correctly (must not be flagged), "novel" once (forbid
# list L2). The "No correction for multiple comparisons" sentence (ledger L3,
# must stay) is already in the base text. Lesson checks: the flask sentence
# is what fake review comment 4 asks to delete (lost facts "catalogue 4411",
# "250 mL"); the replicate range sentence is what the author deletes by hand
# in Word (outside change, lost facts "0.82", "0.91").
PLANTS = [
    (0, 1, " Earlier reports measured this activity in-vitro only."),
    (0, 0, " This is a novel use of the organism."),
    (2, 1, " The same decline was seen in-vitro with purified protein."),
    (0, 2, " The flasks came from a local glassware supplier (catalogue 4411) and each held 250 mL of medium."),
    (2, 2, " Across the replicates the plateau optical density ranged from 0.82 to 0.91."),
]
# (section number, position, old, new): an in-text replacement in the first copy.
REWRITES = [
    (1, 2, "Enzyme activity was measured spectrophotometrically",
     "Enzyme activity was measured in vitro spectrophotometrically"),
]
AUTHOR_COMMENT = ("Growth on the synthetic carbon source", "Check the lag time against the raw data.")
FIXED_TIME = datetime.datetime(2026, 1, 1, 0, 0, 0)

TITLE = "Growth response of a model bacterium to a synthetic carbon source"

# Phrases a Word test can target with find/replace:
#   "the incubation temperature was held at 30 degrees" (occurs exactly once)
#   "cell-free extract" (a real hyphenated compound; first copy of three paragraphs)
SECTIONS = [
    (
        "1. Introduction",
        [
            "Microbial growth on single carbon sources is a common starting point for studies of "
            "central metabolism. A model bacterium was chosen here because its growth is easy to "
            "follow by optical density and because its genetic tools are widely available.",
            "Previous work in this illustrative text suggested that a synthetic carbon source can "
            "support slow growth, although the response varies with medium composition. The aim of "
            "this fictional study was to describe the growth curve and to test whether a "
            "cell-free extract retains the activity of the first enzyme in the pathway.",
            "The remainder of this manuscript is placeholder prose written only to fill several "
            "pages. It contains no real data, and every value below is invented for a software "
            "test of line numbering and review comments.",
        ],
    ),
    (
        "2. Materials and Methods",
        [
            "Cultures were grown in a defined mineral medium in sealed flasks with shaking. "
            "For every run the incubation temperature was held at 30 degrees and the headspace "
            "was refreshed once per day. Optical density was read at 600 nm in a plate reader.",
            "A cell-free extract was prepared by sonication followed by centrifugation, and the "
            "supernatant was kept on ice until use. Protein concentration was estimated with a "
            "dye-binding assay against a standard curve.",
            "Enzyme activity was measured spectrophotometrically by following the change in "
            "absorbance over five minutes. Each assay was repeated three times, and the mean "
            "and standard deviation were reported for every condition.",
            "Statistical comparisons used a two-sided test with a significance threshold of "
            "0.05. No correction for multiple comparisons was applied because only two "
            "conditions were compared in this invented design.",
        ],
    ),
    (
        "3. Results",
        [
            "Growth on the synthetic carbon source began after a lag of roughly one day and "
            "reached a plateau after four days. The final optical density was lower than on the "
            "reference substrate, which is consistent with a slower uptake step.",
            "The cell-free extract showed measurable activity that declined over two hours on "
            "ice. Heat treatment removed the activity completely, which suggests that the signal "
            "depends on a protein component rather than on a small molecule.",
            "Replicate cultures agreed closely during exponential growth and diverged slightly "
            "at the plateau. The spread at the plateau did not change the main conclusion of "
            "this illustrative example.",
        ],
    ),
]

# Repeat the body so the document spans two to three pages.
BODY_REPEATS = 2


def enable_line_numbers(section):
    """Insert w:lnNumType into sectPr at its schema position (after w:pgMar)."""
    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:lnNumType")):
        sect_pr.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    # CT_SectPr order: ... pgSz, pgMar, paperSrc, pgBorders, lnNumType, pgNumType, cols, ...
    successors = ("w:pgNumType", "w:cols", "w:formProt", "w:vAlign", "w:noEndnote",
                  "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
                  "w:printerSettings", "w:sectPrChange")
    for tag in successors:
        nxt = sect_pr.find(qn(tag))
        if nxt is not None:
            nxt.addprevious(ln)
            return
    sect_pr.append(ln)


def set_compatibility_mode(doc, mode):
    """Set w:compatSetting compatibilityMode in word/settings.xml."""
    compat = doc.settings.element.find(qn("w:compat"))
    for cs in compat.findall(qn("w:compatSetting")):
        if cs.get(qn("w:name")) == "compatibilityMode":
            cs.set(qn("w:val"), str(mode))
            return
    raise ValueError("the template has no compatibilityMode setting")


def build(out_path, variant="spike", compat=15):
    doc = Document()
    set_compatibility_mode(doc, compat)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2.0

    cp = doc.core_properties
    cp.author = "Spike Fixture"
    cp.last_modified_by = "Spike Fixture"
    cp.title = TITLE
    cp.created = FIXED_TIME
    cp.modified = FIXED_TIME
    cp.revision = 1

    doc.add_paragraph(TITLE, style="Title")
    comment_run = None
    for sec, (heading, paras) in enumerate(SECTIONS):
        doc.add_paragraph(heading, style="Heading 1")
        for rep in range(BODY_REPEATS):
            for pos, text in enumerate(paras):
                # Keep the target phrases unique: only the first copy carries them.
                if rep > 0:
                    text = (text.replace("the incubation temperature was held at 30 degrees",
                                         "the incubation temperature was kept constant")
                                .replace("cell-free extract", "crude extract"))
                elif variant == "review":
                    for s_, i, old, new in REWRITES:
                        if s_ == sec and i == pos:
                            text = text.replace(old, new)
                    for s_, i, extra in PLANTS:
                        if s_ == sec and i == pos:
                            text = text + extra
                para = doc.add_paragraph(text)
                if variant == "review" and rep == 0 and text.startswith(AUTHOR_COMMENT[0]):
                    comment_run = para.runs[0]
    if comment_run is not None:
        doc.add_comment(comment_run, text=AUTHOR_COMMENT[1], author="Author", initials="AU")

    enable_line_numbers(doc.sections[0])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def main(argv):
    args = argv[1:]
    variant = "spike"
    compat = 15
    if args[:2] == ["--variant", "review"]:
        variant, args = "review", args[2:]
    if args[:1] == ["--compat"]:
        compat, args = int(args[1]), args[2:]
    out = Path(args[0]) if args else (REVIEW_OUT if variant == "review" else DEFAULT_OUT)
    print(build(out, variant, compat))


if __name__ == "__main__":
    main(sys.argv)
