"""Known-answer check for tools/anonymize_docx.py, judged by scripts/verify-anonymized.py.

The expected counts come from the fixture (tests/anon_docx_fixture.py, EXPECTED_DIRTY), not from the
tool. Run: python3 -m unittest tests/test_anonymize_docx.py
"""

import hashlib
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
import zlib
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import anon_docx_fixture as fx  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
TOOL = ROOT / "tools" / "anonymize_docx.py"
REPORTED = ["hidden_text", "protection", "image_uninspected", "embedded"]


def load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def member(path, name):
    with zipfile.ZipFile(path) as z:
        return z.read(name)


class AnonymizeDocx(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.verifier = load(ROOT / "scripts" / "verify-anonymized.py", "verify_anonymized")
        cls.tool = load(TOOL, "anonymize_docx")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.src = fx.build(self.dir / "paper.docx")
        self.out = self.dir / "paper_anonymized.docx"

    def tearDown(self):
        self.tmp.cleanup()

    def run_tool(self, *args):
        r = subprocess.run([sys.executable, str(TOOL), *map(str, args)], capture_output=True, text=True)
        return r.returncode, r.stdout, r.stderr

    def anonymize(self, src=None, out=None, *args):
        if out is None and self.out.exists():
            self.out.unlink()  # the tool refuses an existing output; each run here starts without one
        return self.run_tool("--in", src or self.src, "--out", out or self.out, *args)

    def counts(self, rep):
        return {c: v["count"] for c, v in rep["categories"].items()}

    # ---- the category table
    def test_the_tool_and_the_verifier_know_the_same_categories_with_the_same_defaults(self):
        switchable = {c: v[0] == "clear" for c, v in self.verifier.CATEGORIES.items() if v[0] != "report"}
        self.assertEqual(self.tool.CATEGORIES, switchable)
        self.assertEqual([c for c, v in self.verifier.CATEGORIES.items() if v[0] == "report"], REPORTED)
        self.assertEqual(set(self.verifier.CATEGORIES), set(fx.EXPECTED_DIRTY))

    def test_by_default_every_default_category_is_cleared_and_the_text_is_unchanged(self):
        before = sha(self.src)
        code, out, err = self.anonymize()
        self.assertEqual(code, 0, err)
        rep = self.verifier.verify(self.src, self.out)
        self.assertTrue(self.verifier.passed(rep), rep)
        self.assertEqual([rep["total"], rep["body_equal"], rep["markup_equal"], rep["structure_ok"], rep["unreadable_parts"]],
                         [0, True, True, True, []])
        want = {c: 0 for c in fx.EXPECTED_DIRTY}
        for kept in ["title", "para_ids", *REPORTED]:
            want[kept] = fx.EXPECTED_DIRTY[kept]
        self.assertEqual(self.counts(rep), want)
        status = {c: v["status"] for c, v in rep["categories"].items()}
        for cid, default in self.tool.CATEGORIES.items():
            self.assertEqual(status[cid], "cleared" if default else "kept by option", cid)
        self.assertEqual(rep["not_cleared"], REPORTED)
        self.assertEqual(sha(self.src), before)
        self.assertIn(fx.TRAP, member(self.out, "word/document.xml").decode("utf-8").replace("&quot;", '"'))

    def test_each_category_can_be_kept_on_its_own(self):
        for cid, default in self.tool.CATEGORIES.items():
            if not default:
                continue
            with self.subTest(category=cid):
                code, _, err = self.anonymize(None, None, "--keep", cid)
                self.assertEqual(code, 0, err)
                rep = self.verifier.verify(self.src, self.out)
                got = self.counts(rep)
                self.assertEqual(got[cid], fx.EXPECTED_DIRTY[cid])
                self.assertEqual(rep["total"], fx.EXPECTED_DIRTY[cid])
                self.assertEqual(rep["categories"][cid]["status"], "not cleared")
                self.assertFalse(self.verifier.passed(rep))
                told = self.verifier.verify(self.src, self.out, keep=[cid])
                self.assertEqual(told["categories"][cid]["status"], "kept by option")
                self.assertTrue(self.verifier.passed(told), told)

    def test_title_and_paragraph_ids_are_cleared_only_when_asked(self):
        code, _, err = self.anonymize(None, None, "--also", "title,para_ids")
        self.assertEqual(code, 0, err)
        rep = self.verifier.verify(self.src, self.out, also=["title", "para_ids"])
        self.assertTrue(self.verifier.passed(rep), rep)
        self.assertEqual([self.counts(rep)["title"], self.counts(rep)["para_ids"]], [0, 0])
        doc = member(self.out, "word/document.xml").decode("utf-8")
        self.assertNotIn("paraId", doc)
        self.assertNotIn("textId", doc)
        # The comment's paragraph id is what commentsExtended.xml links to: it stays.
        self.assertIn('w14:paraId="5E6F7A8B"', member(self.out, "word/comments.xml").decode("utf-8"))
        self.assertIn('w15:paraId="5E6F7A8B"', member(self.out, "word/commentsExtended.xml").decode("utf-8"))
        self.assertNotIn(b"An invented manuscript", member(self.out, "docProps/core.xml"))

    def test_dates_become_the_fixed_date_or_are_removed(self):
        self.anonymize()
        for part in ("word/document.xml", "word/comments.xml", "word/footnotes.xml", "word/commentsExtensible.xml"):
            xml = member(self.out, part).decode("utf-8")
            self.assertIn(fx.FIXED_DATE, xml, part)
            self.assertNotIn("2026-01-0", xml.replace(fx.TRAP.replace('"', "&quot;"), "").replace(fx.TRAP, ""), part)
        code, _, err = self.anonymize(None, None, "--dates", "remove")
        self.assertEqual(code, 0, err)
        doc = member(self.out, "word/document.xml").decode("utf-8")
        self.assertEqual(doc.count("w:date="), 1)  # only the one inside the quoted text
        self.assertNotIn("dateUtc", member(self.out, "word/commentsExtensible.xml").decode("utf-8"))
        self.assertTrue(self.verifier.passed(self.verifier.verify(self.src, self.out)))
        core = member(self.out, "docProps/core.xml").decode("utf-8")
        self.assertEqual(core.count(fx.FIXED_DATE), 3)
        self.assertIn("<cp:revision>1</cp:revision>", core)

    def test_the_report_says_what_was_changed_and_removed(self):
        code, out, err = self.anonymize()
        # One line: the daemon's Python runner reads the last line of stdout as JSON.
        self.assertEqual(len(out.strip().splitlines()), 1)
        report = json.loads(out)
        self.assertEqual(report["output"], str(self.out))
        verifier_cli = subprocess.run(
            [sys.executable, str(ROOT / "scripts" / "verify-anonymized.py"), "--before", str(self.src), "--after", str(self.out), "--json"],
            capture_output=True, text=True,
        )
        self.assertEqual(len(verifier_cli.stdout.strip().splitlines()), 1)
        self.assertEqual(json.loads(verifier_cli.stdout)["total"], 0)
        ch = report["changed"]
        self.assertEqual(ch["word/document.xml"], {"author": 3, "date": 3, "rsid": 4})
        self.assertEqual(ch["word/comments.xml"], {"author": 1, "initials": 1, "date": 1})
        self.assertEqual(ch["docProps/custom.xml"], {"property": 1})
        self.assertEqual(ch["word/people.xml"], {"person": 2, "presenceInfo": 2})
        self.assertEqual(ch["word/settings.xml"], {"rsid": 3, "docVar": 2, "attachedTemplate": 1, "docId": 2})
        self.assertEqual(ch["word/_rels/settings.xml.rels"], {"attachedTemplate": 1})
        self.assertEqual(ch["docProps/app.xml"], {"Company": 1, "Manager": 1, "HyperlinkBase": 1})
        self.assertEqual(ch["word/media/image1.jpeg"], {"segments": 2})
        self.assertEqual(ch["word/media/image2.png"], {"segments": 2})
        self.assertEqual(report["removed"], ["customXml/_rels/item1.xml.rels", "customXml/item1.xml",
                                             "customXml/itemProps1.xml", "docProps/thumbnail.jpeg"])
        self.assertEqual(report["kept"], ["para_ids", "title"])

    def test_the_source_names_are_gone_from_every_part(self):
        self.anonymize()
        with zipfile.ZipFile(self.out) as z:
            blob = b"".join(z.read(n) for n in z.namelist())
        for name in (b"Test Author", b"Second Person", b"Example Lab", b"Lab Head", b"example.org", b"EX-0",
                     b"test.author", b"RefManagerUser", b"00EF3456", b"Lab.dotm", b"1A2B3C4D"):
            self.assertNotIn(name, blob)
        self.assertIn(b"An invented manuscript", blob)  # the title is kept by default
        # app.xml: time, template and version are left alone.
        app = member(self.out, "docProps/app.xml")
        for kept in (b"<TotalTime>135</TotalTime>", b"<Template>Normal.dotm</Template>", b"<AppVersion>16.0000</AppVersion>"):
            self.assertIn(kept, app)

    def test_removed_parts_leave_no_reference_and_the_rest_keeps_its_place(self):
        self.anonymize()
        gone = {"customXml/_rels/item1.xml.rels", "customXml/item1.xml", "customXml/itemProps1.xml", "docProps/thumbnail.jpeg"}
        with zipfile.ZipFile(self.src) as a, zipfile.ZipFile(self.out) as b:
            self.assertEqual([n for n in a.namelist() if n not in gone], b.namelist())
            self.assertIsNone(b.testzip())
            for name in ("word/styles.xml", "customXml/item2.xml", "customXml/itemProps2.xml", "word/media/image3.tiff",
                         "word/embeddings/oleObject1.bin", "word/commentsExtended.xml"):
                self.assertEqual(a.read(name), b.read(name), name)
            for name in b.namelist():
                if name.endswith((".xml", ".rels")):
                    ET.fromstring(b.read(name))
            doc = b.read("word/document.xml").decode("utf-8")
            rels = b.read("word/_rels/document.xml.rels").decode("utf-8")
            types = b.read("[Content_Types].xml").decode("utf-8")
            root_rels = b.read("_rels/.rels").decode("utf-8")
        self.assertNotIn("item1.xml", rels)
        self.assertIn("../customXml/item2.xml", rels)
        self.assertIn("https://journal.example/guide", rels)
        self.assertNotIn("itemProps1", types)
        self.assertIn("itemProps2", types)
        self.assertNotIn("thumbnail", root_rels)
        self.assertIn("word/document.xml", root_rels)
        # Namespace declarations that mc:Ignorable names must survive, or Word asks to repair the file.
        self.assertIn('xmlns:w14="%s"' % fx.W14, doc)
        self.assertIn('xmlns:w15="%s"' % fx.W15, doc)
        self.assertIn('mc:Ignorable="w14 w15"', doc)
        self.assertIn('w14:paraId="0A1B2C3D"', doc)
        self.assertTrue(doc.startswith('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'))
        settings = member(self.out, "word/settings.xml").decode("utf-8")
        self.assertIn("<w:documentProtection", settings)
        self.assertIn("compatibilityMode", settings)

    # ---- the zip container
    def test_zip_comments_extra_fields_and_member_times_are_cleared(self):
        before = sha(self.src)
        with zipfile.ZipFile(self.src) as z:  # the fixture really carries them
            self.assertEqual(z.comment, fx.ZIP_COMMENT)
            self.assertEqual(z.getinfo("word/document.xml").comment, fx.ENTRY_COMMENT)
            self.assertEqual(z.getinfo("word/document.xml").extra, fx.ZIP_EXTRA)
        code, out, err = self.anonymize()
        self.assertEqual(code, 0, err)
        with zipfile.ZipFile(self.out) as z:
            self.assertEqual(z.comment, b"")
            for info in z.infolist():
                self.assertEqual((info.comment, info.extra, info.date_time), (b"", b"", fx.ZIP_EPOCH), info.filename)
        blob = Path(self.out).read_bytes()
        for leak in (b"test.author", b"example.org", fx.ZIP_COMMENT, fx.ENTRY_COMMENT, fx.ZIP_EXTRA):
            self.assertNotIn(leak, blob)
        self.assertEqual(json.loads(out)["zip"], {"archive_comment": 1, "entry_comment": 1, "extra_field": 1, "timestamp": 2})
        rep = self.verifier.verify(self.src, self.out)
        self.assertEqual([rep["categories"]["zip_meta"]["count"], rep["body_equal"]], [0, True])
        self.assertTrue(self.verifier.passed(rep))
        self.assertEqual(sha(self.src), before)
        # Asked to keep them, the tool copies them as they are.
        self.anonymize(None, None, "--keep", "zip_meta")
        with zipfile.ZipFile(self.out) as z:
            d = z.getinfo("word/document.xml")
            self.assertEqual((z.comment, d.comment, d.extra, d.date_time), (fx.ZIP_COMMENT, fx.ENTRY_COMMENT, fx.ZIP_EXTRA, (2026, 1, 2, 3, 4, 6)))

    def test_only_the_comments_as_an_archiver_would_leave_them(self):
        # A docx without any other mark, with an archive comment and a member comment.
        clean = fx.build(self.dir / "clean.docx", **fx.clean_marks())
        src = fx.rezip(clean, self.dir / "commented.docx", comment=b"Created by someone@example.org",
                       entry_comment=b"Author: someone@example.org")
        self.assertEqual(self.verifier.verify(src, src)["categories"]["zip_meta"]["count"], 2)
        code, out, err = self.anonymize(src)
        self.assertEqual(code, 0, err)
        self.assertNotIn(b"someone@example.org", Path(self.out).read_bytes())
        self.assertEqual(json.loads(out)["zip"], {"archive_comment": 1, "entry_comment": 1})
        self.assertTrue(self.verifier.passed(self.verifier.verify(src, self.out)))
        # A docx without comments is judged as before: nothing to clear, nothing reported.
        code, out, err = self.anonymize(clean)
        self.assertEqual([code, json.loads(out)["zip"]], [0, {}], err)
        self.assertTrue(self.verifier.passed(self.verifier.verify(clean, self.out)))

    # ---- pictures
    def test_picture_metadata_is_dropped_and_the_picture_data_is_not_touched(self):
        self.anonymize()
        jpg = member(self.out, "word/media/image1.jpeg")
        self.assertEqual(jpg, b"\xff\xd8" + fx.JFIF + fx.ICC + fx.SCAN)
        png = member(self.out, "word/media/image2.png")
        self.assertEqual(png, fx.png(text=False))
        # Every chunk that is left has a valid checksum.
        pos = 8
        while pos < len(png):
            n = int.from_bytes(png[pos : pos + 4], "big")
            self.assertEqual(zlib.crc32(png[pos + 4 : pos + 8 + n]) & 0xFFFFFFFF, int.from_bytes(png[pos + 8 + n : pos + 12 + n], "big"))
            pos += 12 + n

    def test_exif_orientation_is_read_in_both_byte_orders(self):
        for order in ("<", ">"):
            for orientation in (1, 3, 6, 8):
                self.assertEqual(self.tool.exif_orientation(fx.exif(orientation, order)[6:]), orientation)
        self.assertEqual(self.tool.exif_orientation(b"II*\x00\x08\x00\x00\x00\x00\x00"), None)
        self.assertEqual(self.tool.exif_orientation(b"garbage"), -1)
        for data, state, removed in (
            (fx.jpeg(), "stripped", 2),
            (fx.jpeg(order=">"), "stripped", 2),
            (fx.jpeg(with_exif=False, comment=b""), "none", 0),
            (fx.jpeg(orientation=6), "rotated", 0),
            (b"\xff\xd8 broken", "unreadable", 0),
            (b"not a jpeg", "unreadable", 0),
            (fx.jpeg()[: -len(fx.SCAN)], "unreadable", 0),  # cut before the scan
        ):
            new, n, got = self.tool.strip_jpeg(data)
            self.assertEqual((got, n), (state, removed))
            if state != "stripped":
                self.assertEqual(new, data)

    def test_a_rotated_picture_is_left_alone_and_the_verifier_says_so(self):
        src = fx.build(self.dir / "rot.docx", extra={"word/media/image9.jpeg": fx.jpeg(orientation=6)})
        code, out, err = self.anonymize(src)
        self.assertEqual(code, 0, err)
        self.assertEqual(member(self.out, "word/media/image9.jpeg"), fx.jpeg(orientation=6))
        self.assertEqual(json.loads(out)["changed"]["word/media/image9.jpeg"], {"left_rotated": 1})
        rep = self.verifier.verify(src, self.out)
        self.assertEqual((rep["categories"]["image_meta"]["count"], rep["categories"]["image_meta"]["status"]), (1, "not cleared"))
        self.assertFalse(self.verifier.passed(rep))
        self.assertTrue(self.verifier.passed(self.verifier.verify(src, self.out, keep=["image_meta"])))

    # ---- a document made by a library, reopened by it
    def test_a_real_document_still_opens_and_reads_the_same(self):
        from docx import Document

        src = ROOT / "tests" / "fixtures" / "review_manuscript.docx"
        for extra in ((), ("--dates", "remove"), ("--also", "title,para_ids")):
            code, _, err = self.anonymize(src, None, *extra)
            self.assertEqual(code, 0, err)
            self.assertEqual([p.text for p in Document(self.out).paragraphs], [p.text for p in Document(src).paragraphs])
            self.assertEqual(Document(self.out).core_properties.author, "Anonymous")
            also = ["title", "para_ids"] if "--also" in extra else []
            rep = self.verifier.verify(src, self.out, also=also)
            self.assertTrue(self.verifier.passed(rep), rep)

    # ---- as before
    def test_a_chosen_name_is_written_and_escaped(self):
        code, _, err = self.anonymize(None, None, "--author", 'R&D "Copy" <1>', "--initials", "R&D")
        self.assertEqual(code, 0, err)
        rep = self.verifier.verify(self.src, self.out, 'R&D "Copy" <1>', "R&D")
        self.assertEqual(rep["total"], 0, rep["remaining"])
        self.assertEqual(self.verifier.verify(self.src, self.out)["total"] > 0, True)

    def test_another_prefix_and_single_quotes_are_handled(self):
        for kw in (dict(prefix="x"), dict(quote="'")):
            with self.subTest(**kw):
                src = fx.build(self.dir / "odd.docx", **kw)
                out = self.dir / ("odd_out_%s.docx" % "".join(kw))
                code, _, err = self.anonymize(src, out)
                self.assertEqual(code, 0, err)
                rep = self.verifier.verify(src, out)
                self.assertEqual([rep["total"], rep["body_equal"], rep["markup_equal"]], [0, True, True], rep)

    def test_a_docx_without_the_optional_parts_works(self):
        gone = {"word/people.xml", "word/comments.xml", "docProps/custom.xml", "docProps/app.xml", "word/footnotes.xml",
                "word/settings.xml", "customXml/item1.xml", "docProps/thumbnail.jpeg"}
        src = fx.build(self.dir / "small.docx", parts=gone)
        code, _, err = self.anonymize(src)
        self.assertEqual(code, 0, err)
        rep = self.verifier.verify(src, self.out)
        self.assertEqual([rep["total"], rep["structure_ok"]], [0, True], rep)

    def test_running_it_on_its_own_output_changes_nothing_more(self):
        self.anonymize()
        again = self.dir / "again.docx"
        code, out, _ = self.anonymize(self.out, again)
        self.assertEqual(code, 0)
        self.assertEqual([json.loads(out)["changed"], json.loads(out)["removed"]], [{}, []])
        with zipfile.ZipFile(self.out) as a, zipfile.ZipFile(again) as b:
            self.assertEqual([a.read(n) for n in a.namelist()], [b.read(n) for n in b.namelist()])

    def test_an_unknown_category_is_refused(self):
        for flag in ("--keep", "--also"):
            code, _, err = self.anonymize(None, None, flag, "hidden_text")
            self.assertEqual(code, 2)
            self.assertIn("unknown category", err)
        self.assertFalse(self.out.exists())

    def test_the_output_may_not_be_the_source_under_any_name(self):
        before = sha(self.src)
        link = self.dir / "link.docx"
        os.symlink(self.src, link)
        hard = self.dir / "hard.docx"
        os.link(self.src, hard)
        sub = self.dir / "sub"
        sub.mkdir()
        spellings = [self.src, link, hard, sub / ".." / "paper.docx", Path(str(self.src).replace("paper.docx", "./paper.docx"))]
        for out in spellings:
            with self.subTest(out=str(out)):
                code, _, err = self.anonymize(self.src, out)
                self.assertEqual(code, 2)
                self.assertIn("same file", err)
                self.assertEqual(sha(self.src), before)
        # And the other way round: the source given through a link.
        self.assertEqual(self.anonymize(link, self.src)[0], 2)
        self.assertEqual(sha(self.src), before)
        self.assertEqual(sorted(p.name for p in self.dir.iterdir()), ["hard.docx", "link.docx", "paper.docx", "sub"])

    def test_no_half_written_file_is_left_behind(self):
        bad = self.dir / "bad.docx"
        bad.write_text("not a zip")
        code, _, err = self.anonymize(bad)
        self.assertEqual(code, 2)
        self.assertFalse(self.out.exists())
        missing_dir = self.dir / "nowhere" / "out.docx"
        self.assertEqual(self.anonymize(self.src, missing_dir)[0], 2)
        self.anonymize()
        self.assertEqual(sorted(p.name for p in self.dir.iterdir()), ["bad.docx", "paper.docx", "paper_anonymized.docx"])

    def test_an_existing_output_is_refused_unless_overwrite_is_given(self):
        self.out.write_text("old")
        code, _, err = self.run_tool("--in", self.src, "--out", self.out)
        self.assertEqual(code, 2)
        self.assertIn("already exists", err)
        self.assertEqual(self.out.read_text(), "old")
        # A dangling link is an existing name too.
        dangling = self.dir / "dangling.docx"
        os.symlink(self.dir / "nowhere.docx", dangling)
        self.assertEqual(self.run_tool("--in", self.src, "--out", dangling)[0], 2)
        self.assertFalse((self.dir / "nowhere.docx").exists())
        code, _, err = self.run_tool("--in", self.src, "--out", self.out, "--overwrite")
        self.assertEqual(code, 0, err)
        self.assertTrue(zipfile.is_zipfile(self.out))
        self.assertEqual(sorted(p.name for p in self.dir.iterdir()), ["dangling.docx", "paper.docx", "paper_anonymized.docx"])

    def test_a_file_that_appears_while_the_copy_is_written_is_not_replaced(self):
        # The first check passes (no file), then the name is taken before the copy is placed.
        tmp = self.dir / "work.tmp"
        tmp.write_bytes(b"new copy")
        self.out.write_text("theirs")
        with self.assertRaises(self.tool.Refused):
            self.tool.place(str(tmp), str(self.out))
        self.assertEqual(self.out.read_text(), "theirs")
        self.assertTrue(tmp.exists())
        free = self.dir / "free.docx"
        self.tool.place(str(tmp), str(free))
        self.assertEqual([free.read_bytes(), tmp.exists()], [b"new copy", False])


if __name__ == "__main__":
    unittest.main()
