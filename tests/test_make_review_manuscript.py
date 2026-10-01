"""Known-answer check for scripts/make-test-manuscript.py --variant review.

Counts are the ones the stage-2 acceptance test relies on (ledger L1-L3 and
the one comment by the author). Run: python3 -m unittest tests/test_make_review_manuscript.py
"""

import importlib.util
import re
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location("mtm", ROOT / "scripts" / "make-test-manuscript.py")
mtm = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(mtm)


class ReviewVariant(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = Path(cls.tmp.name) / "review.docx"
        mtm.build(cls.path, "review")
        cls.paras = [p.text for p in Document(cls.path).paragraphs]
        cls.body = "\n".join(cls.paras)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def count(self, pattern, flags=0):
        return len(re.findall(pattern, self.body, flags))

    def test_ledger_plants(self):
        self.assertEqual(self.count(r"\bin-(?:vitro|vivo|situ)\b", re.I), 2)
        self.assertEqual(self.count(r"\bin vitro\b"), 1)
        self.assertEqual(self.count(r"\bnovel(?:ty)?\b", re.I), 1)
        self.assertEqual(self.count(r"No correction for multiple comparisons was applied"), 2)

    def test_lesson_plants(self):
        # Comment 4 deletes the flask sentence (lost facts); the author deletes
        # the replicate range sentence by hand (outside change).
        self.assertEqual(self.count(r"\(catalogue 4411\) and each held 250 mL"), 1)
        self.assertEqual(self.count(r"ranged from 0\.82 to 0\.91"), 1)
        self.assertEqual(self.count(r"antifoam", re.I), 0)

    def test_review_targets(self):
        self.assertEqual(self.count(r"the incubation temperature was held at 30 degrees"), 1)
        self.assertEqual(self.count(r"cell-free extract"), 3)
        self.assertEqual(self.count(r"A cell-free extract was prepared"), 1)
        self.assertEqual(self.count(r"reached a plateau after four days"), 2)

    def test_one_author_comment(self):
        with zipfile.ZipFile(self.path) as z:
            xml = z.read("word/comments.xml").decode("utf-8")
        self.assertEqual(xml.count("<w:comment "), 1)
        self.assertIn('w:author="Author"', xml)
        self.assertNotRegex(xml, r"\[R\d+\.\d+\]")

    def test_spike_variant_unchanged(self):
        spike = Path(self.tmp.name) / "spike.docx"
        mtm.build(spike)
        text = "\n".join(p.text for p in Document(spike).paragraphs)
        self.assertNotIn("in-vitro", text)
        self.assertNotIn("novel", text)


def compat_mode(path):
    """compatibilityMode value stored in word/settings.xml, or None."""
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/settings.xml").decode("utf-8")
    m = re.search(r'<w:compatSetting\b[^>]*w:name="compatibilityMode"[^>]*>', xml)
    return re.search(r'w:val="(\d+)"', m.group(0)).group(1) if m else None


class CompatibilityMode(unittest.TestCase):
    """Word refuses Comment.resolved in a document whose compatibilityMode is
    14 (Word 2010), the value in the python-docx default template. The
    builder writes 15 so the fixtures can show a resolved rationale comment."""

    def test_both_variants_are_mode_15(self):
        with tempfile.TemporaryDirectory() as tmp:
            for variant in ("spike", "review"):
                out = mtm.build(Path(tmp) / f"{variant}.docx", variant)
                self.assertEqual(compat_mode(out), "15", variant)

    def test_committed_fixtures_are_mode_15(self):
        for name in ("spike_manuscript.docx", "review_manuscript.docx"):
            self.assertEqual(compat_mode(ROOT / "tests" / "fixtures" / name), "15", name)

    def test_compat_14_is_only_on_request(self):
        # The acceptance run builds a mode-14 copy in its run folder to show
        # that a refresh still reports its counts when Word refuses the resolve.
        with tempfile.TemporaryDirectory() as tmp:
            out = mtm.build(Path(tmp) / "old.docx", "review", compat=14)
            self.assertEqual(compat_mode(out), "14")
            mtm.main(["x", "--variant", "review", "--compat", "14", str(Path(tmp) / "cli.docx")])
            self.assertEqual(compat_mode(Path(tmp) / "cli.docx"), "14")
            # Same text as the mode-15 build.
            new = mtm.build(Path(tmp) / "new.docx", "review")
            self.assertEqual(
                [p.text for p in Document(out).paragraphs],
                [p.text for p in Document(new).paragraphs],
            )


if __name__ == "__main__":
    unittest.main()
