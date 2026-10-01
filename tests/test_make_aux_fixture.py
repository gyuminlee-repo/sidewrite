"""The fake manuscript of the stage 4 acceptance run (scripts/e2e/make_aux_fixture.py)."""
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "e2e" / "make_aux_fixture.py"
spec = importlib.util.spec_from_file_location("make_aux_fixture", SCRIPT)
fx = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fx)


def run(args):
    return subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True)


class AuxFixture(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        d = Path(cls.tmp.name)
        cls.current = fx.build(d / "current.docx", "current")
        cls.v1 = fx.build(d / "v1.docx", "v1")
        cls.dir = d

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_abstract_is_over_the_limit_and_under_its_heading(self):
        paras = Document(self.current).paragraphs
        at = [p.text for p in paras].index("Abstract")
        self.assertEqual(paras[at].style.name, "Heading 1")
        self.assertGreater(len(paras[at + 1].text.split()), fx.ABSTRACT_LIMIT)
        self.assertEqual(paras[at + 2].text, fx.KEYWORDS)
        self.assertEqual(len(fx.KEYWORDS.split(":")[1].split(",")), 3)

    def test_each_planted_format_violation_is_there_once(self):
        doc = Document(self.current)
        plain_p = italic_species = plain_species = plain_formula = sub_digit = 0
        flush = indented = 0
        for p in doc.paragraphs:
            if p.style.name != "Normal" or len(p.text.split()) < 20:
                continue
            first = p.paragraph_format.first_line_indent
            if first and first.pt > 0:
                indented += 1
            else:
                flush += 1
            for r in p.runs:
                if "(p < 0.05)" in r.text and not r.italic:
                    plain_p += 1
                if fx.SPECIES in r.text:
                    if r.italic:
                        italic_species += 1
                    else:
                        plain_species += 1
                if fx.FORMULA in r.text:
                    plain_formula += 1
                if r.font.subscript and r.text == "2":
                    sub_digit += 1
        self.assertEqual(
            (plain_p, plain_species, plain_formula, flush),
            (fx.PLANTED["stat"], fx.PLANTED["species"], fx.PLANTED["subscript"], fx.PLANTED["indent"]),
        )
        self.assertEqual((italic_species, sub_digit), (1, 1))
        # The flush paragraph is an outlier of a clear majority.
        self.assertGreaterEqual(indented / (indented + flush), 0.8)
        self.assertEqual(len(doc.tables), fx.PLANTED["empty_table"])
        self.assertTrue(all(not c.text for row in doc.tables[0].rows for c in row.cells))

    def test_the_earlier_version_differs_in_two_sentences_and_one_paragraph(self):
        now = [p.text for p in Document(self.current).paragraphs]
        old = [p.text for p in Document(self.v1).paragraphs]
        self.assertEqual(len(now), len(old) + 1)
        changed = [t for t in old if t not in now]
        self.assertEqual(len(changed), 2)
        self.assertEqual([t for t in now if t not in old and t != "".join(x for x, _ in fx.LAST)].__len__(), 2)

    def test_the_marked_tool_takes_the_pair_in_both_modes(self):
        for mode in ("track_changes", "color"):
            out = self.dir / ("marked_%s.docx" % mode)
            r = run([ROOT / "tools" / "build_marked_diff.py", "--baseline", self.v1, "--clean",
                     self.current, "--out", out, "--json", "--mode", mode])
            self.assertEqual(r.returncode, 0, r.stderr[-500:])
            report = json.loads(r.stdout.strip().splitlines()[-1])
            self.assertEqual((report["failures"], report["written"]), (0, True))
            self.assertEqual((report["paragraphs"]["n_mod"], report["paragraphs"]["n_ins"]), (2, 1))

    def test_the_anonymized_copy_passes_its_verifier(self):
        out = self.dir / "anon.docx"
        r = run([ROOT / "tools" / "anonymize_docx.py", "--in", self.current, "--out", out])
        self.assertEqual(r.returncode, 0, r.stderr[-500:])
        v = run([ROOT / "scripts" / "verify-anonymized.py", "--before", self.current, "--after", out, "--json"])
        self.assertEqual(v.returncode, 0, v.stderr[-500:])
        rep = json.loads(v.stdout.strip().splitlines()[-1])
        self.assertEqual((rep["total"], rep["body_equal"]), (0, True))

    def test_the_committed_fixtures_read_like_a_fresh_build(self):
        for variant, built in (("current", self.current), ("v1", self.v1)):
            committed = fx.OUT[variant]
            self.assertTrue(committed.is_file(), committed)
            self.assertEqual(
                [p.text for p in Document(committed).paragraphs],
                [p.text for p in Document(built).paragraphs],
            )

    def test_profile_and_dump(self):
        r = run([SCRIPT, "--profile"])
        self.assertEqual(json.loads(r.stdout), fx.PROFILE)
        rows = fx.dump(self.current)
        self.assertEqual(sum(1 for x in rows if x["table_level"]), 4)
        self.assertEqual([x["text"] for x in rows if x["heading"]], ["Abstract", "Methods", "Results"])


if __name__ == "__main__":
    unittest.main()
