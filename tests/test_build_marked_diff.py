"""Known-answer check for tools/build_marked_diff.py.

Two small documents (all text invented): one word replaced, one sentence
added, one paragraph removed, one paragraph added, with an italic run and a
subscript run in text that does not change. Run:
python3 -m unittest tests/test_build_marked_diff.py
"""

import hashlib
import re
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

ROOT = Path(__file__).resolve().parent.parent
TOOL = ROOT / "tools" / "build_marked_diff.py"

TITLE = "Growth of an invented strain"
REMOVED = "A sentence that will be removed entirely from the text of this document."
ADDED_SENTENCE = "A second sentence was added here."
ADDED_PARA = "Wholly new closing words follow here without any overlap."


def write_doc(path, days, extra_sentence, removed, added, red_word=None):
    d = Document()
    d.add_paragraph(TITLE)
    p = d.add_paragraph()
    p.add_run("The strain ")
    p.add_run("Escherichia coli").italic = True
    p.add_run(" grew on CO")
    p.add_run("2").font.subscript = True
    p.add_run(" at 30 degrees for %s days in every flask." % days)
    p = d.add_paragraph("The yield was low in every flask of the series.")
    if extra_sentence:
        p.add_run(" " + ADDED_SENTENCE)
    if removed:
        d.add_paragraph(REMOVED)
    if added:
        d.add_paragraph(ADDED_PARA)
    if red_word:
        r = d.add_paragraph("A closing remark with a ").add_run(red_word)
        r.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    d.save(path)
    return path


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def doc_xml(path):
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode("utf-8")


def runs(path):
    """[(text, colour or None, struck)] of every run, paragraph by paragraph."""
    out = []
    for p in Document(path).paragraphs:
        out.append([(r.text, str(r.font.color.rgb) if r.font.color.rgb else None, bool(r.font.strike)) for r in p.runs])
    return out


def texts(path):
    return [p.text for p in Document(path).paragraphs if p.text]


class BuildMarkedDiff(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.base = write_doc(self.dir / "baseline.docx", "two", False, True, False)
        self.clean = write_doc(self.dir / "clean.docx", "three", True, False, True)
        self.out = self.dir / "marked.docx"

    def tearDown(self):
        self.tmp.cleanup()

    def run_tool(self, *args, baseline=None, clean=None, out=None):
        cmd = [sys.executable, str(TOOL), "--baseline", str(baseline or self.base), "--clean", str(clean or self.clean),
               "--out", str(out or self.out), *args]
        r = subprocess.run(cmd, capture_output=True, text=True, cwd=self.dir)
        return r.returncode, r.stdout, r.stderr

    def test_insertions_are_coloured_and_deletions_struck_in_the_given_colours(self):
        code, out, err = self.run_tool("--insert-rgb", "FF0000", "--delete-rgb", "808080")
        self.assertEqual(code, 0, out + err)
        self.assertRegex(out, r"VERIFY SUMMARY: checks \d+, failures 0")
        flat = [r for para in runs(self.out) for r in para]
        inserted = "".join(t for t, colour, struck in flat if colour == "FF0000" and not struck)
        deleted = "".join(t for t, colour, struck in flat if struck)
        self.assertIn("three", inserted)
        self.assertIn(ADDED_SENTENCE, inserted)
        self.assertIn(ADDED_PARA, inserted)
        self.assertNotIn("strain", inserted)
        self.assertIn("two", deleted)
        self.assertIn(REMOVED, deleted)
        self.assertNotIn("yield", deleted)
        # Every struck run has the delete colour and nothing else has it.
        self.assertTrue(all(colour == "808080" for _, colour, struck in flat if struck))
        self.assertFalse(any(colour == "808080" and not struck for _, colour, struck in flat))

    def test_accepting_gives_the_clean_text_and_rejecting_gives_the_baseline_text(self):
        self.run_tool("--insert-rgb", "FF0000", "--delete-rgb", "808080")
        marked = runs(self.out)
        accept = ["".join(t for t, _, struck in para if not struck) for para in marked]
        reject = ["".join(t for t, colour, struck in para if struck or colour != "FF0000") for para in marked]
        self.assertEqual([t for t in accept if t], texts(self.clean))
        self.assertEqual([t for t in reject if t], texts(self.base))

    def test_italic_and_subscript_runs_are_as_many_as_in_the_clean_file(self):
        self.run_tool()
        marked, clean = doc_xml(self.out), doc_xml(self.clean)
        for mark in ("<w:i/>", "<w:vertAlign "):
            self.assertEqual(marked.count(mark), clean.count(mark), mark)
            self.assertGreater(clean.count(mark), 0)

    def test_only_the_document_part_differs_from_the_clean_file(self):
        before = sha(self.clean), sha(self.base)
        self.run_tool()
        self.assertEqual((sha(self.clean), sha(self.base)), before)
        with zipfile.ZipFile(self.clean) as a, zipfile.ZipFile(self.out) as b:
            self.assertEqual(a.namelist(), b.namelist())
            for name in a.namelist():
                if name != "word/document.xml":
                    self.assertEqual(a.read(name), b.read(name), name)
        self.assertEqual(sorted(p.name for p in self.dir.iterdir()), ["baseline.docx", "clean.docx", "marked.docx"])

    def test_other_colours_come_from_the_command_line_only(self):
        code, out, err = self.run_tool("--insert-rgb", "0000ff", "--delete-rgb", "00AA00")
        self.assertEqual(code, 0, out + err)
        xml = doc_xml(self.out)
        self.assertIn('w:val="0000FF"', xml)
        self.assertIn('w:val="00AA00"', xml)
        self.assertNotIn('w:val="FF0000"', xml)
        for bad in ("red", "FF00", "GG0000"):
            code, _, err = self.run_tool("--insert-rgb", bad)
            self.assertEqual(code, 2, bad)

    def test_track_changes_mode_writes_revision_marks_under_the_given_author(self):
        code, out, err = self.run_tool("--mode", "track_changes", "--author", "Reviewer copy")
        self.assertEqual(code, 0, out + err)
        xml = doc_xml(self.out)
        self.assertGreater(len(re.findall(r"<w:ins ", xml)), 0)
        self.assertGreater(len(re.findall(r"<w:del ", xml)), 0)
        self.assertEqual(set(re.findall(r'w:author="([^"]*)"', xml)), {"Reviewer copy"})
        self.assertNotIn("<w:strike", xml)

    def test_the_output_may_not_be_one_of_the_inputs(self):
        before = sha(self.clean), sha(self.base)
        for out in (self.clean, self.base, self.dir / "." / "clean.docx"):
            code, _, err = self.run_tool(out=out)
            self.assertEqual(code, 2)
            self.assertIn("same file", err)
        for report in (self.clean, self.base, self.out):
            code, _, err = self.run_tool("--report", str(report))
            self.assertEqual(code, 2)
            self.assertIn("same file", err)
        self.assertFalse(self.out.exists())
        self.assertEqual((sha(self.clean), sha(self.base)), before)

    def test_the_report_has_counts_and_no_document_text(self):
        report = self.dir / "report.json"
        code, out, err = self.run_tool("--report", str(report))
        self.assertEqual(code, 0, out + err)
        text = report.read_text(encoding="utf-8")
        self.assertIn('"failures": 0', text)
        for words in ("invented strain", "three", ADDED_SENTENCE, REMOVED):
            self.assertNotIn(words, text)
            self.assertNotIn(words, out)

    def test_a_clean_file_that_already_uses_the_insert_colour_is_refused_with_a_way_out(self):
        clean = write_doc(self.dir / "red.docx", "three", True, False, True, red_word="warning")
        code, _, err = self.run_tool(clean=clean)
        self.assertEqual(code, 2)
        self.assertIn("--insert-rgb", err)
        self.assertFalse(self.out.exists())
        code, out, err = self.run_tool("--insert-rgb", "0000FF", clean=clean)
        self.assertEqual(code, 0, out + err)

    def test_a_missing_or_broken_input_is_an_error_and_writes_nothing(self):
        self.assertEqual(self.run_tool(baseline=self.dir / "none.docx")[0], 2)
        bad = self.dir / "bad.docx"
        bad.write_text("not a zip")
        self.assertEqual(self.run_tool(clean=bad)[0], 2)
        self.assertFalse(self.out.exists())

    def test_identical_files_give_a_copy_without_marks(self):
        code, out, err = self.run_tool(baseline=self.clean)
        self.assertEqual(code, 0, out + err)
        self.assertEqual(doc_xml(self.out).count("<w:strike"), 0)
        self.assertEqual(texts(self.out), texts(self.clean))

    def test_the_tool_carries_no_manuscript_specific_value(self):
        src = TOOL.read_text(encoding="utf-8")
        for word in ("/mnt/", "/home/", "BITE", "jrp_config", "submission.json", "extorquens", "4EA72E", "JRP_"):
            self.assertNotIn(word, src, word)
        # A whole word: the WordprocessingML namespace ends in "...ingml".
        self.assertNotRegex(src, r"\bgml\b")
        self.assertNotRegex(src, r"\bR[12]\b")
        self.assertNotRegex(src, r"20\d\d-\d\d-\d\d")
        self.assertNotRegex(src, r"\b2026\d{4}\b")


if __name__ == "__main__":
    unittest.main()
