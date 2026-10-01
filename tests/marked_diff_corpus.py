"""Invented baseline/clean pairs for the marked-diff tool, and an own reading of its output.

Nothing here uses the tool's code. A document is a model (a list of blocks); the same model gives the docx file
and the structure the tool's output must show: after every mark is accepted the structure of the clean model,
after every mark is rejected the structure of the baseline model. Structure means every paragraph (empty ones
too) with its text, tabs and typed breaks and with what it holds besides text (a field with its code and its
result, a picture with the bytes it shows, an equation with its text, a symbol, a note reference), every table
with its rows, cells and cell paragraphs, and the text inside every bookmark.

A paragraph is {"k": "p", "id": ..., "pieces": [...]}. A piece is a string or a tuple:
("tab",), ("br", "line" | "page" | "column" | "all" | "none"), ("link", text), ("bm", name, text), ("b", text),
("fn",) and ("fn2",) (references to two notes), ("pic",) and ("pic2",) (two inline pictures that differ), ("field", code, result), ("cite", number,
shown text) (a field as a reference manager writes it: a long code cut into several runs), ("simple", code,
result) (a w:fldSimple), ("fopen", code) and ("fend",) (the two ends of a field that runs over several
paragraphs, as a bibliography does), ("sym",), ("nbh",) (a hyphen that does not break), ("shy",) (a soft
hyphen), ("cr",), ("math", text), ("cm", text) (text with a comment on it), ("sdt", text) (text in a content
control), ("ins", text) and ("del", text) (a change that is still pending in the input file).
A table is {"k": "t", "id": ..., "rows": [[cell, ...], ...]}; a cell is a list of piece lists (its paragraphs).
A table may have "merge": [("h", row, first column, last column) | ("v", column, first row, last row)]; the
cells a merge takes in must be empty in the model.
{"k": "bs", "name": ..., "n": ...} and {"k": "be", "n": ...} are the start and the end of a bookmark that stand
between paragraphs, as Word writes them when a bookmark covers whole paragraphs.
"""

import copy
import hashlib
import io
import json
import random
import struct
import zipfile
import zlib

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt

INSERT_RGB = "FF0000"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
BIBLIOGRAPHY = 'ADDIN ZOTERO_BIBL {"uncited":[],"omitted":[],"custom":[]} CSL_BIBLIOGRAPHY'


def _png(shade=0):
    """A one-pixel picture; another shade is another picture."""
    chunk = lambda kind, data: struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)  # noqa: E731
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 0, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(bytes([0, shade]))) + chunk(b"IEND", b""))


def cite_code(number, shown):
    """The code of a citation field as a reference manager writes it: its name and one long JSON value (invented)."""
    item = {"id": number, "uris": ["http://example.invalid/items/INV%04d" % number],
            "itemData": {"id": number, "type": "article-journal",
                         "title": "An invented title about the %s, number %d" % (sentence(number)[4:-1], number),
                         "container-title": "Journal of Invented Results", "volume": str(10 + number),
                         "page": "%d-%d" % (100 * number, 100 * number + 9),
                         "author": [{"family": "Invented", "given": "A."}, {"family": "Madeup", "given": "B."}],
                         "issued": {"date-parts": [[2031]]}}}
    value = {"citationID": "inv%d" % number, "properties": {"formattedCitation": shown, "plainCitation": shown,
                                                            "noteIndex": 0},
             "citationItems": [item], "schema": "https://example.invalid/schema/csl-citation.json"}
    return " ADDIN ZOTERO_ITEM CSL_CITATION %s " % json.dumps(value, separators=(",", ":"))


PAGE, COLUMN, LINE_ALL = "<page>", "<column>", "<line:all>"

_A = ["quiet", "narrow", "distant", "early", "plain", "heavy", "bright", "silent"]
_N = ["harbour", "orchard", "bridge", "lantern", "meadow", "ferry", "quarry", "garden"]
_V = ["shelters", "follows", "crosses", "carries", "watches", "divides", "gathers", "borders"]
_O = ["river", "valley", "station", "market", "tower", "village", "forest", "island"]
_T = ["before the rain", "in the long winter", "near the old road", "without any hurry", "during the harvest"]


def sentence(i):
    """A fixed invented sentence; different numbers share few words."""
    return "The %s %s %s the %s %s number %d." % (
        _A[i % 8], _N[(i * 3) % 8], _V[(i * 5) % 8], _O[(i * 7 + i // 8) % 8], _T[i % 5], i)


def text(*numbers):
    return " ".join(sentence(i) for i in numbers)


def P(ident, *pieces):
    return {"k": "p", "id": ident, "pieces": list(pieces)}


def T(ident, rows):
    return {"k": "t", "id": ident, "rows": rows}


def base_model():
    return [
        P("h1", "Opening notes"),
        P("a", text(0, 1, 2)),
        P("b", text(3, 4)),
        P("e", ),
        P("c", text(5, 6, 7)),
        T("t1", [[[[text(30)]], [[text(31)]]], [[[text(32)]], [[text(33)], [text(34)]]]]),
        P("d", text(8, 9)),
        P("h2", "Closing notes"),
        P("f", text(10, 11, 12)),
        P("g", text(13, 14)),
    ]


# ---------------------------------------------------------------- the model as a docx file
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), v)
    return e


def _run(text_):
    r = _el("w:r")
    t = _el("w:t")
    t.text = text_
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    return r


_IDS = iter(range(500, 100000))


def _fill(par, pieces):
    """Fill a python-docx paragraph with the pieces of the model."""
    p = par._p
    for piece in pieces:
        if isinstance(piece, str):
            p.append(_run(piece))
            continue
        kind = piece[0]
        if kind == "tab":
            r = _el("w:r")
            r.append(_el("w:tab"))
            p.append(r)
        elif kind == "br":
            r = _el("w:r")
            attrs = {"line": {}, "page": {"type": "page"}, "column": {"type": "column"},
                     "all": {"clear": "all"}, "none": {"clear": "none"}}[piece[1]]
            r.append(_el("w:br", **attrs))
            p.append(r)
        elif kind == "link":
            link = _el("w:hyperlink", anchor="somewhere")
            link.append(_run(piece[1]))
            p.append(link)
        elif kind == "bm":
            n = str(next(_IDS))
            p.append(_el("w:bookmarkStart", id=n, name=piece[1]))
            p.append(_run(piece[2]))
            p.append(_el("w:bookmarkEnd", id=n))
        elif kind == "b":
            r = _run(piece[1])
            rpr = _el("w:rPr")
            rpr.append(_el("w:b"))
            r.insert(0, rpr)
            p.append(r)
        elif kind in ("fn", "fn2"):
            r = _el("w:r")
            r.append(_el("w:footnoteReference", id="2" if kind == "fn" else "3"))
            p.append(r)
        elif kind in ("pic", "pic2"):
            par.add_run().add_picture(io.BytesIO(_png(0 if kind == "pic" else 200)), width=Pt(24), height=Pt(24))
        elif kind in ("field", "cite", "fopen", "fend"):
            code = cite_code(piece[1], piece[2]) if kind == "cite" else " %s " % piece[1] if kind != "fend" else ""
            # a long code is cut into several runs, as Word does
            chunks = [code[i:i + 90] for i in range(0, len(code), 90)]
            parts = {"fend": ["end"], "fopen": ["begin"] + chunks + ["separate"]}.get(
                kind, ["begin"] + chunks + ["separate", None, "end"])
            for part in parts:
                r = _el("w:r")
                if part in ("begin", "separate", "end"):
                    r.append(_el("w:fldChar", fldCharType=part))
                elif part is None:
                    r = _run(piece[2])
                else:
                    text_ = _el("w:instrText")
                    text_.text = part
                    text_.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    r.append(text_)
                p.append(r)
        elif kind == "simple":
            fld = _el("w:fldSimple", instr=" %s " % piece[1])
            fld.append(_run(piece[2]))
            p.append(fld)
        elif kind in ("sym", "nbh", "shy", "cr"):
            r = _el("w:r")
            r.append(_el("w:sym", font="Symbol", char="F061") if kind == "sym" else
                     _el({"nbh": "w:noBreakHyphen", "shy": "w:softHyphen", "cr": "w:cr"}[kind]))
            p.append(r)
        elif kind == "math":
            p.append(parse_xml('<m:oMath xmlns:m="%s"><m:r><m:t>%s</m:t></m:r></m:oMath>' % (M_NS, piece[1])))
        elif kind == "cm":
            run = par.add_run(piece[1])
            par.part.document.add_comment(run, text="An invented remark.", author="Somebody", initials="S")
        elif kind == "sdt":
            sdt = _el("w:sdt")
            pr = _el("w:sdtPr")
            pr.append(_el("w:alias", val="Invented"))
            body = _el("w:sdtContent")
            body.append(_run(piece[1]))
            sdt.append(pr)
            sdt.append(body)
            p.append(sdt)
        elif kind in ("ins", "del"):
            w = _el("w:" + kind, id=str(next(_IDS)), author="Somebody", date="2031-01-02T03:04:05Z")
            r = _run(piece[1])
            if kind == "del":
                r[0].tag = qn("w:delText")
            w.append(r)
            p.append(w)
        else:
            raise ValueError(kind)


FOOTNOTES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:footnote w:type="separator" w:id="0"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:type="continuationSeparator" w:id="1"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    '<w:footnote w:id="2"><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:footnote>'
    '<w:footnote w:id="3"><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:footnote>'
    "</w:footnotes>"
)
NOTE_TEXT = {"fn": "An invented note at the foot of the page.", "fn2": "Another invented note."}


def write(model, path, header=None):
    """Save the model as a docx file. With a ("fn",) piece a footnotes part is added."""
    d = Document()
    for block in model:
        if block["k"] == "p":
            _fill(d.add_paragraph(), block["pieces"])
        elif block["k"] in ("bs", "be"):
            attrs = {"id": str(block["n"]), "name": block["name"]} if block["k"] == "bs" else {"id": str(block["n"])}
            d.element.body.find(qn("w:sectPr")).addprevious(_el("w:bookmarkStart" if block["k"] == "bs" else "w:bookmarkEnd", **attrs))
        else:
            rows = block["rows"]
            table = d.add_table(rows=len(rows), cols=len(rows[0]))
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    tc = table.cell(i, j)
                    _fill(tc.paragraphs[0], cell[0])
                    for pieces in cell[1:]:
                        _fill(tc.add_paragraph(), pieces)
            for kind, a, b, c in block.get("merge", []):
                if kind == "h":
                    table.cell(a, b).merge(table.cell(a, c))
                else:
                    table.cell(b, a).merge(table.cell(c, a))
    if header:
        d.sections[0].header.paragraphs[0].text = header
    d.save(path)
    used = {piece[0] for b in model if b["k"] == "p" for piece in b["pieces"] if piece in (("fn",), ("fn2",))}
    if used:
        _add_footnotes(path, used)
    return path


def _add_footnotes(path, used):
    """Add the footnotes part, with the notes the model refers to."""
    notes = FOOTNOTES % (NOTE_TEXT["fn"], NOTE_TEXT["fn2"])
    for kind, n in (("fn", 2), ("fn2", 3)):
        if kind not in used:
            notes = notes.replace('<w:footnote w:id="%d"><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:footnote>' % (n, NOTE_TEXT[kind]), "")
    with zipfile.ZipFile(path) as z:
        parts = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in parts:
            if info.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-'
                    b'officedocument.wordprocessingml.footnotes+xml"/></Types>')
            if info.filename == "word/_rels/document.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdNotes" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    b'relationships/footnotes" Target="footnotes.xml"/></Relationships>')
            z.writestr(info.filename, data)
        z.writestr("word/footnotes.xml", notes)


# ---------------------------------------------------------------- the structure a model stands for
def _text(pieces):
    out = ""
    for piece in pieces:
        if isinstance(piece, str):
            out += piece
        elif piece[0] == "tab":
            out += "\t"
        elif piece[0] == "br":
            out += {"line": "\n", "none": "\n", "page": PAGE, "column": COLUMN, "all": LINE_ALL}[piece[1]]
        elif piece[0] in ("link", "b", "ins", "cm", "sdt"):
            out += piece[1]
        elif piece[0] in ("bm", "field", "cite", "simple"):
            out += piece[2]
    return out


def _inline(pieces):
    """What the pieces hold besides text, in order: ("field", code, result), ("drawing", digest of the picture),
    ("oMath", text), ("sym", font, character), (kind,) for the rest."""
    out = []
    for piece in pieces:
        kind = "" if isinstance(piece, str) else piece[0]
        if kind in ("pic", "pic2"):
            out.append(("drawing", hashlib.sha256(_png(0 if kind == "pic" else 200)).hexdigest()))
        elif kind in ("field", "simple"):
            out.append(("field", " %s " % piece[1], piece[2]))
        elif kind == "cite":
            out.append(("field", cite_code(piece[1], piece[2]), piece[2]))
        elif kind == "fopen":
            out.append(("field-open", " %s " % piece[1]))
        elif kind == "fend":
            out.append(("fldChar", "end"))
        elif kind == "math":
            out.append(("oMath", piece[1]))
        elif kind == "sym":
            out.append(("sym", "Symbol", "F061"))
        elif kind in ("fn", "fn2"):
            out.append(("footnoteReference", NOTE_TEXT[kind]))
        elif kind in ("nbh", "shy", "cr"):
            out.append(({"nbh": "noBreakHyphen", "shy": "softHyphen", "cr": "cr"}[kind],))
    return tuple(out)


def _comments(pieces):
    return sum(1 for piece in pieces if not isinstance(piece, str) and piece[0] == "cm")


def _layout(block):
    """(grid columns, rows of cells as (model column, span, vertical merge)) of a table of the model."""
    rows = [[{"col": j, "span": 1, "vm": None} for j in range(len(row))] for row in block["rows"]]
    for kind, a, b, c in block.get("merge", []):
        if kind == "h":
            rows[a] = [x for x in rows[a] if not b < x["col"] <= c]
            next(x for x in rows[a] if x["col"] == b)["span"] = c - b + 1
        else:
            for i in range(b, c + 1):
                next(x for x in rows[i] if x["col"] == a)["vm"] = "restart" if i == b else "continue"
    return len(block["rows"][0]), rows


def structure_of(model):
    """(items in order, {bookmark name: text}). An item is ("p", text), followed by ("x", what it holds besides text,
    see _inline) when there is something, and by ("c", (comment range starts, ends, references)) when it has
    comments; or ("t", rows of cells of paragraph texts), followed by ("g", (grid columns, rows of cells as
    (span, vertical merge)))."""
    items, marks, opened = [], {}, {}
    for block in model:
        if block["k"] == "p":
            items.append(("p", _text(block["pieces"])))
            if _inline(block["pieces"]):
                items.append(("x", _inline(block["pieces"])))
            if _comments(block["pieces"]):
                items.append(("c", (_comments(block["pieces"]),) * 3))
            marks.update({piece[1]: piece[2] for piece in block["pieces"] if not isinstance(piece, str) and piece[0] == "bm"})
            for name in opened.values():
                marks[name] += _text(block["pieces"])
        elif block["k"] == "bs":
            opened[block["n"]] = block["name"]
            marks[block["name"]] = ""
        elif block["k"] == "be":
            opened.pop(block["n"])
        else:
            grid, layout = _layout(block)
            items.append(("t", tuple(tuple(tuple(_text(p) for p in block["rows"][i][x["col"]]) for x in row)
                                     for i, row in enumerate(layout))))
            items.append(("g", (grid, tuple(tuple((x["span"], x["vm"]) for x in row) for row in layout))))
    return items, marks


# ---------------------------------------------------------------- an own reading of a marked file
def _on(el):
    return el is not None and el.get(qn("w:val")) not in ("0", "false", "off")


def _rpr_state(rpr):
    if rpr is None:
        return None
    if _on(rpr.find(qn("w:strike"))):
        return "del"
    colour = rpr.find(qn("w:color"))
    return "ins" if colour is not None and (colour.get(qn("w:val")) or "").upper() == INSERT_RGB else None


def _mark_state(rpr):
    if rpr is None:
        return None
    return "ins" if rpr.find(qn("w:ins")) is not None else "del" if rpr.find(qn("w:del")) is not None else None


def view(path, mode, side):
    """The body of a marked file with every mark accepted (side "accept") or rejected ("reject").

    Track changes: runs in w:ins / w:del, rows with w:trPr/w:ins|w:del, paragraph marks with w:pPr/w:rPr/w:ins|w:del.
    Colour: struck runs are deletions and runs in the insert colour insertions; a paragraph mark formatted that
    way is a deleted or inserted paragraph mark; a row whose paragraph marks are all deleted (or all inserted)
    is a deleted (inserted) row. Dropping a paragraph mark joins the paragraph with the one after it."""
    doc = Document(path)
    body = doc.element.body
    for blip in body.iter(A_BLIP):                    # a picture is the bytes it shows, not the id that points to them
        if blip.get(R_EMBED) in doc.part.related_parts:
            blip.set(R_EMBED, hashlib.sha256(doc.part.related_parts[blip.get(R_EMBED)].blob).hexdigest())
    notes = {}                                        # a note reference is the text of its note, not its id
    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/footnotes"):
            root = parse_xml(rel.target_part.blob)
            notes = {n.get(qn("w:id")): "".join(t.text or "" for t in n.iter(qn("w:t"))) for n in root}
    for ref in body.iter(qn("w:footnoteReference")):
        ref.set("note", notes.get(ref.get(qn("w:id")), "no such note"))
    drop = "del" if side == "accept" else "ins"
    track = mode == "track_changes"

    def para_state(p):
        rpr = p.find(qn("w:pPr") + "/" + qn("w:rPr"))
        return _mark_state(rpr) if track else _rpr_state(rpr)

    def row_state(tr):
        if track:
            return _mark_state(tr.find(qn("w:trPr")))
        states = {para_state(p) for p in tr.iter(qn("w:p"))}
        return states.pop() if len(states) == 1 else None

    def run_state(r):
        if track:
            tags = [a.tag for a in r.iterancestors()]
            return "del" if qn("w:del") in tags else "ins" if qn("w:ins") in tags else None
        return _rpr_state(r.find(qn("w:rPr")))

    for tr in list(body.iter(qn("w:tr"))):
        if row_state(tr) == drop:
            tr.getparent().remove(tr)
    for tbl in list(body.iter(qn("w:tbl"))):
        if tbl.find(qn("w:tr")) is None:
            tbl.getparent().remove(tbl)
    for r in list(body.iter(qn("w:r"))):
        if run_state(r) == drop:
            r.getparent().remove(r)
    for p in list(body.iter(qn("w:p"))):
        if para_state(p) != drop:
            continue
        nxt, between = p.getnext(), []
        while nxt is not None and nxt.tag in (qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            between.append(nxt)  # a bookmark that starts or ends between two paragraphs does not keep them apart
            nxt = nxt.getnext()
        if nxt is None or nxt.tag != qn("w:p"):
            continue  # stays: the comparison then shows a paragraph too many
        at = 1 if len(nxt) and nxt[0].tag == qn("w:pPr") else 0
        for child in [c for c in p if c.tag != qn("w:pPr")] + between:  # the markers keep their place in the text
            nxt.insert(at, child)
            at += 1
        p.getparent().remove(p)
    return body


def _para_text(p):
    return "".join(_run_text(r) for r in p.iter(qn("w:r")))


def _run_text(r):
    out = ""
    for ch in r:
        if ch.tag in (qn("w:t"), qn("w:delText")):
            out += ch.text or ""
        elif ch.tag == qn("w:tab"):
            out += "\t"
        elif ch.tag == qn("w:br"):
            kind, clear = ch.get(qn("w:type")), ch.get(qn("w:clear"))
            if kind == "page":
                out += PAGE
            elif kind == "column":
                out += COLUMN
            elif clear in (None, "none"):
                out += "\n"
            else:
                out += "<line:%s>" % clear
    return out


def _cell_layout(tc):
    """(columns the cell spans, None | "restart" | "continue") of a w:tc."""
    span = tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan"))
    vm = tc.find(qn("w:tcPr") + "/" + qn("w:vMerge"))
    return (int(span.get(qn("w:val"))) if span is not None else 1,
            None if vm is None else vm.get(qn("w:val")) or "continue")


def _contents(p):
    """What a paragraph holds besides text, read from its XML in the terms of _inline. A field that does not end
    in the paragraph is ("field-open", code); a field mark without its begin is ("fldChar", kind)."""
    out, field = [], None
    for e in p.iter():
        if not isinstance(e.tag, str):
            continue
        if e.tag == qn("w:fldChar"):
            kind = e.get(qn("w:fldCharType"))
            if kind == "begin" and field is None:
                field = {"depth": 1, "code": "", "result": "", "sep": False}
            elif field is None:
                out.append(("fldChar", kind))
            elif kind == "begin":
                field["depth"] += 1
            elif kind == "separate" and field["depth"] == 1:
                field["sep"] = True
            elif kind == "end":
                field["depth"] -= 1
                if field["depth"] == 0:
                    out.append(("field", field["code"], field["result"]))
                    field = None
        elif e.tag in (qn("w:instrText"), qn("w:delInstrText")):
            if field is not None:
                field["code"] += e.text or ""
        elif e.tag in (qn("w:t"), qn("w:delText")):
            if field is not None and field["sep"]:
                field["result"] += e.text or ""
        elif e.tag == qn("w:fldSimple"):
            out.append(("field", e.get(qn("w:instr")), "".join(t.text or "" for t in e.iter(qn("w:t")))))
        elif e.tag == qn("w:drawing"):
            out.append(("drawing", next((b.get(R_EMBED) for b in e.iter(A_BLIP)), None)))
        elif e.tag == "{%s}oMath" % M_NS:
            out.append(("oMath", "".join(t.text or "" for t in e.iter("{%s}t" % M_NS))))
        elif e.tag == qn("w:sym"):
            out.append(("sym", e.get(qn("w:font")), e.get(qn("w:char"))))
        elif e.tag == qn("w:footnoteReference"):
            out.append(("footnoteReference", e.get("note")))
        elif e.tag in [qn("w:" + k) for k in ("endnoteReference", "noBreakHyphen", "softHyphen", "cr", "pict", "object")]:
            out.append((e.tag.split("}")[-1],))
    if field is not None:
        out.append(("field-open", field["code"]))
    return tuple(out)


def structure_read(body):
    """The same two values as structure_of, read from a body element."""
    items, marks, opened = [], {}, {}
    for child in body:
        if child.tag == qn("w:p"):
            items.append(("p", _para_text(child)))
            if _contents(child):
                items.append(("x", _contents(child)))
            notes = tuple(len(list(child.iter(qn("w:" + k)))) for k in ("commentRangeStart", "commentRangeEnd", "commentReference"))
            if any(notes):
                items.append(("c", notes))
        elif child.tag == qn("w:tbl"):
            items.append(("t", tuple(
                tuple(tuple(_para_text(p) for p in tc.findall(qn("w:p"))) for tc in tr.findall(qn("w:tc")))
                for tr in child.findall(qn("w:tr")))))
            items.append(("g", (len(child.findall(qn("w:tblGrid") + "/" + qn("w:gridCol"))), tuple(
                tuple(_cell_layout(tc) for tc in tr.findall(qn("w:tc"))) for tr in child.findall(qn("w:tr"))))))
    for e in body.iter(qn("w:bookmarkStart"), qn("w:bookmarkEnd"), qn("w:r")):
        if e.tag == qn("w:bookmarkStart"):
            opened[e.get(qn("w:id"))] = e.get(qn("w:name"))
            marks.setdefault(e.get(qn("w:name")), "")
        elif e.tag == qn("w:bookmarkEnd"):
            opened.pop(e.get(qn("w:id")), None)
        else:
            piece = _run_text(e)
            for name in opened.values():
                marks[name] += piece
    marks.pop("_GoBack", None)
    return items, marks


def without_page_breaks(items):
    """Colour mode shows a deleted page or column break as nothing (a struck break would still break the page)."""
    strip = lambda s: s.replace(PAGE, "").replace(COLUMN, "")
    return [("p", strip(x[1])) if x[0] == "p" else
            ("t", tuple(tuple(tuple(strip(p) for p in cell) for cell in row) for row in x[1])) if x[0] == "t" else x
            for x in items]


def split_reject(items):
    """(body paragraphs, tables): where a table stands among the paragraphs is not compared on the reject side, and
    neither are comments (a comment is no tracked change: the output has the clean document's)."""
    return [x for x in items if x[0] in ("p", "x")], [x for x in items if x[0] in ("t", "g")]


# ---------------------------------------------------------------- changes
def _at(model, ident):
    return next(i for i, b in enumerate(model) if b["id"] == ident)


def _reword(pieces, n=3, new="altered"):
    """Change one word of the first string piece."""
    for k, piece in enumerate(pieces):
        if isinstance(piece, str) and len(piece.split(" ")) > n:
            words = piece.split(" ")
            words[n] = new
            pieces[k] = " ".join(words)
            return
    raise ValueError("no text to change")


def reword(ident, n=3, new="altered"):
    return lambda m: _reword(m[_at(m, ident)]["pieces"], n, new)


def add_para(before, numbers, ident="new"):
    def op(m):
        new = P(ident, text(*numbers)) if numbers else P(ident, ("bm", "mark1", "Figure 3"), " " + text(6, 7))
        m.insert(len(m) if before is None else _at(m, before), new)
    return op


def del_para(ident):
    return lambda m: m.pop(_at(m, ident))


def split_para(ident):
    def op(m):
        i = _at(m, ident)
        first, rest = m[i]["pieces"][0].split(". ", 1)
        m[i]["pieces"][0] = first + "."
        m.insert(i + 1, P(ident + "2", rest))
    return op


def merge_paras(ident):
    def op(m):
        i = _at(m, ident)
        j = next(k for k in range(i + 1, len(m)) if m[k]["k"] == "p")
        m[i]["pieces"] = [_text(m[i]["pieces"]) + " " + _text(m[j]["pieces"])]
        m.pop(j)
    return op


def add_empty(before):
    return lambda m: m.insert(len(m) if before is None else _at(m, before), P("empty2"))


def replace_para(ident, numbers):
    def op(m):
        m[_at(m, ident)]["pieces"] = [text(*numbers)]
    return op


def move_para(ident, before):
    def op(m):
        block = m.pop(_at(m, ident))
        m.insert(_at(m, before), block)
    return op


def set_pieces(ident, *pieces):
    def op(m):
        m[_at(m, ident)]["pieces"] = list(pieces)
    return op


def table(ident, change):
    return lambda m: change(m[_at(m, ident)]["rows"])


def add_row(rows):
    rows.append([[[text(40 + j)]] for j in range(len(rows[0]))])


def del_row(rows):
    rows.pop()


def first_row_gone(rows):
    rows.pop(0)


def add_col(rows):
    for i, row in enumerate(rows):
        row.append([[text(44 + i)]])


def del_col(rows):
    for row in rows:
        row.pop()


def cell_reword(rows):
    _reword(rows[0][0][0])


def cell_add_para(rows):
    rows[0][1].append([text(46)])


def cell_del_para(rows):
    rows[1][1].pop()


def cell_del_first_para(rows):
    rows[1][1].pop(0)


def cell_replace(rows):
    rows[1][0][0] = [text(47)]


def add_table(before, ident="t2"):
    def op(m):
        m.insert(_at(m, before), T(ident, [[[[text(50)]], [[text(51)]]]]))
    return op


def move_sentence(src, dst):
    """Take the last sentence of one paragraph to the front of another."""
    def op(m):
        head, tail = m[_at(m, src)]["pieces"][0].rsplit(". ", 1)
        m[_at(m, src)]["pieces"][0] = head + "."
        m[_at(m, dst)]["pieces"][0] = tail + " " + m[_at(m, dst)]["pieces"][0]
    return op


def rewrite_keeping(ident, keep, numbers):
    """Rewrite a paragraph: one of its sentences stays, between new ones."""
    def op(m):
        old = m[_at(m, ident)]["pieces"][0].split(". ")
        kept = old[keep] if old[keep].endswith(".") else old[keep] + "."
        m[_at(m, ident)]["pieces"] = [text(numbers[0]) + " " + kept + " " + text(*numbers[1:])]
    return op


def only(*idents):
    """Drop every block but these."""
    def op(m):
        m[:] = [b for b in m if b["id"] in idents]
    return op


def to_front(ident):
    def op(m):
        m.insert(0, m.pop(_at(m, ident)))
    return op


def to_end(ident):
    def op(m):
        m.append(m.pop(_at(m, ident)))
    return op


def span(first, last, name="whole"):
    """A bookmark over whole paragraphs: its start before the block `first`, its end after the block `last`."""
    def op(m):
        m.insert(_at(m, last) + 1, {"k": "be", "id": name + "-end", "n": 77})
        m.insert(_at(m, first), {"k": "bs", "id": name + "-start", "name": name, "n": 77})
    return op


def merge(ident, *merges, empty=()):
    """Merge cells of a table; `empty` are the (row, column) cells the merges take in, which lose their text."""
    def op(m):
        block = m[_at(m, ident)]
        for i, j in empty:
            block["rows"][i][j] = [[]]
        block["merge"] = list(merges)
    return op


def unmerge(ident):
    return lambda m: m[_at(m, ident)].pop("merge")


def both(*ops):
    def op(m):
        for o in ops:
            o(m)
    return op


def nothing(m):
    return None


def variant(*ops):
    """A baseline model: the base with these changes."""
    def make():
        m = base_model()
        for o in ops:
            o(m)
        return m
    return make


BR = lambda kind: set_pieces("c", text(5, 6), ("br", kind), text(7))  # noqa: E731
WITH_TAB = set_pieces("c", text(5), ("tab",), text(6, 7))
WITH_LINK = set_pieces("c", text(5) + " ", ("link", "a linked phrase"), " " + text(6, 7))
WITH_MARK = set_pieces("c", text(5) + " ", ("bm", "mark1", "Figure 3"), " " + text(6, 7))
WITH_BOLD = set_pieces("c", text(5) + " ", ("b", "a bold phrase"), " " + text(6, 7))
WITH_NOTE = set_pieces("d", text(8), ("fn",), " " + text(9))
TWO_TABLES = add_table("h2")
MERGED_ACROSS = merge("t1", ("h", 0, 0, 1), empty=[(0, 1)])
MERGED_DOWN = merge("t1", ("v", 0, 0, 1), empty=[(1, 0)])
WITH_PICTURE = set_pieces("c", text(5) + " ", ("pic",), " " + text(6, 7))
WITH_FIELD = set_pieces("c", text(5) + " ", ("field", "PAGE", "3"), " " + text(6, 7))
PICTURE_PARA = lambda m: m.insert(_at(m, "b"), P("pic", text(20) + " ", ("pic",)))  # noqa: E731
PICTURE_ONLY = lambda m: m.insert(_at(m, "b"), P("pic", ("pic",)))  # noqa: E731
FIELD_PARA = lambda m: m.insert(_at(m, "b"), P("fld", text(20) + " ", ("field", "PAGE", "3"), "."))  # noqa: E731
CITE1, CITE2, CITE3 = ("cite", 1, "[1]"), ("cite", 2, "[2]"), ("cite", 3, "[3]")
WITH_CITES = set_pieces("c", text(5) + " ", CITE1, " " + text(6) + " ", CITE2, " " + text(7))
CAPTION = lambda m: m.insert(_at(m, "d"), P("cap", "Figure ", ("field", "SEQ Figure \\* ARABIC", "1"), ". " + text(40) + " See ", ("field", "REF _Ref4242 \\h", "Table 1"), "."))  # noqa: E731
CAPTION2 = lambda m: m.insert(_at(m, "d"), P("cap", "Figure ", ("field", "SEQ Figure \\* ARABIC", "1"), ". " + text(40) + " See ", ("field", "REF _Ref4242 \\h", "Table 2"), "."))  # noqa: E731
WITH_SIGNS = set_pieces("c", text(5) + " well", ("nbh",), "known ", ("sym",), " " + text(6), ("cr",), text(7), ("shy",))
WITH_MATH = set_pieces("c", text(5) + " ", ("math", "x+y"), " " + text(6, 7))
WITH_COMMENT = set_pieces("c", text(5) + " ", ("cm", "a remarked phrase"), " " + text(6, 7))
WITH_SIMPLE = set_pieces("c", text(5) + " ", ("simple", "PAGE", "3"), " " + text(6, 7))
WITH_CONTROL = set_pieces("c", text(5) + " ", ("sdt", "a controlled phrase"), " " + text(6, 7))
BIB = lambda m: m.extend([P("bib1", ("fopen", BIBLIOGRAPHY), "1. " + text(41)), P("bib2", "2. " + text(42)), P("bib3", "3. " + text(43), ("fend",))])  # noqa: E731
CITED_PARA = lambda m: m.insert(_at(m, "b"), P("cited", text(20) + " ", CITE3, " " + text(21)))  # noqa: E731
SPLIT_AT_MARK = set_pieces("c", text(5))
PENDING_INS = set_pieces("b", text(3), ("ins", " " + text(4)))
PENDING_DEL = set_pieces("b", text(3, 4), ("del", " " + text(20)))

# name -> (baseline model maker, change that gives the clean model, refusal pattern or None)
CASES = {
    "no change": (variant(), nothing, None),
    "a word changed": (variant(), reword("a"), None),
    "a word changed in two paragraphs": (variant(), both(reword("a"), reword("f")), None),
    "a paragraph added at the start": (variant(), add_para("h1", (20, 21)), None),
    "a paragraph added in the middle": (variant(), add_para("b", (20, 21)), None),
    "a paragraph added before a table": (variant(), add_para("t1", (20, 21)), None),
    "a paragraph added after a table": (variant(), add_para("d", (20, 21)), None),
    "a paragraph added at the end": (variant(), add_para(None, (20, 21)), None),
    "two paragraphs added at the end": (variant(), both(add_para(None, (20, 21)), add_para(None, (22, 23), "new2")), None),
    "a paragraph deleted at the start": (variant(), del_para("h1"), None),
    "a paragraph deleted in the middle": (variant(), del_para("b"), None),
    "a paragraph deleted before a table": (variant(), del_para("c"), None),
    "a paragraph deleted after a table": (variant(), del_para("d"), None),
    "a paragraph deleted at the end": (variant(), del_para("g"), None),
    "two paragraphs deleted at the end": (variant(), both(del_para("g"), del_para("f")), None),
    "a paragraph split in two": (variant(), split_para("a"), None),
    "the last paragraph split in two": (variant(), split_para("g"), None),
    "two paragraphs merged": (variant(), merge_paras("a"), None),
    "the last two paragraphs merged": (variant(), merge_paras("f"), None),
    "a heading merged into the next paragraph": (variant(), merge_paras("h2"), None),
    "an empty paragraph added": (variant(), add_empty("f"), None),
    "an empty paragraph added at the end": (variant(), add_empty(None), None),
    "an empty paragraph removed": (variant(), del_para("e"), None),
    "a paragraph replaced": (variant(), replace_para("b", (22, 23)), None),
    "the last paragraph replaced": (variant(), replace_para("g", (22, 23)), None),
    "a paragraph moved": (variant(), move_para("f", "a"), None),
    "a table row added": (variant(), table("t1", add_row), None),
    "a table row removed": (variant(), table("t1", del_row), None),
    "the first table row removed": (variant(), table("t1", first_row_gone), None),
    "a table column added": (variant(), table("t1", add_col), None),
    "a table column removed": (variant(), table("t1", del_col), None),
    "a word changed in a cell": (variant(), table("t1", cell_reword), None),
    "a paragraph added in a cell": (variant(), table("t1", cell_add_para), None),
    "the last paragraph of a cell deleted": (variant(), table("t1", cell_del_para), None),
    "the first paragraph of a cell deleted": (variant(), table("t1", cell_del_first_para), None),
    "the text of a cell replaced": (variant(), table("t1", cell_replace), None),
    "a table added": (variant(), add_table("h2"), None),
    "the second of two tables removed": (variant(TWO_TABLES), del_para("t2"), None),
    "the first of two tables removed": (variant(TWO_TABLES), del_para("t1"), None),
    "a line break in a changed paragraph": (variant(BR("line")), reword("c"), None),
    "a page break in a changed paragraph": (variant(BR("page")), reword("c"), None),
    "a column break in a changed paragraph": (variant(BR("column")), reword("c"), None),
    "a break that clears floats in a changed paragraph": (variant(BR("all")), reword("c"), None),
    "a break written with clear none": (variant(BR("none")), reword("c"), None),
    "a break written with clear none, nothing changed": (variant(BR("none")), nothing, None),
    "a page break added": (variant(), BR("page"), None),
    "a page break removed": (variant(BR("page")), set_pieces("c", text(5, 6, 7)), None),
    "a column break removed": (variant(BR("column")), set_pieces("c", text(5, 6, 7)), None),
    "a line break removed": (variant(BR("line")), set_pieces("c", text(5, 6) + " " + text(7)), None),
    "a page break turned into a column break": (variant(BR("page")), BR("column"), None),
    "a paragraph with a page break deleted": (variant(BR("page")), del_para("c"), None),
    "a tab in a changed paragraph": (variant(WITH_TAB), reword("c"), None),
    "a link in a changed paragraph": (variant(WITH_LINK), reword("c"), None),
    "bold text in a changed paragraph": (variant(WITH_BOLD), reword("c"), None),
    "a bookmark in a changed paragraph": (variant(WITH_MARK), both(WITH_MARK, reword("c")), None),
    "a bookmark in a paragraph that is split": (variant(), both(SPLIT_AT_MARK, add_para("t1", (), "c2")), None),
    "a pending insertion in the clean file": (variant(), both(PENDING_INS, reword("a")), None),
    "a pending deletion in the clean file": (variant(), both(PENDING_DEL, reword("a")), None),
    "a pending insertion in the baseline": (variant(PENDING_INS), reword("b"), None),
    "a footnote and a header, change elsewhere": (variant(WITH_NOTE), both(WITH_NOTE, reword("a")), None),
    "a footnote in the changed paragraph": (variant(WITH_NOTE), both(WITH_NOTE, reword("d")), None),
    "a footnote in a deleted paragraph": (variant(WITH_NOTE), del_para("d"), None),
    "a footnote taken out, another stays": (
        variant(WITH_NOTE, set_pieces("f", text(10), ("fn2",), " " + text(11, 12))), set_pieces("d", text(8) + " " + text(9)), None),
    "a footnote swapped for another": (variant(WITH_NOTE), set_pieces("d", text(8), ("fn2",), " " + text(9)), None),
    "a footnote added": (variant(), WITH_NOTE, None),
    "a cell merged across two columns": (variant(), MERGED_ACROSS, None),
    "a cell merged down a column": (variant(), MERGED_DOWN, None),
    "a merged cell split again": (variant(MERGED_DOWN), both(unmerge("t1"), table("t1", lambda rows: rows[1].__setitem__(0, [[text(49)]]))), None),
    "a word changed in a table with a merged cell": (variant(MERGED_DOWN), table("t1", lambda rows: _reword(rows[0][1][0])), None),
    "a merge across turned into a merge down": (variant(MERGED_ACROSS), both(unmerge("t1"), MERGED_DOWN, table("t1", lambda rows: rows[0].__setitem__(1, [[text(31)]]))), None),
    "a paragraph with a picture added": (variant(), PICTURE_PARA, None),
    "a paragraph that is only a picture added": (variant(), PICTURE_ONLY, None),
    "a paragraph with a field added": (variant(), FIELD_PARA, None),
    "a picture in an unchanged paragraph, change elsewhere": (variant(WITH_PICTURE), reword("a"), None),
    "a field in an unchanged paragraph, change elsewhere": (variant(WITH_FIELD), reword("a"), None),
    "a picture added to a paragraph": (variant(), WITH_PICTURE, None),
    "a word changed in a paragraph with a picture": (variant(WITH_PICTURE), reword("c"), None),
    "a word changed in a paragraph with a field": (variant(WITH_FIELD), reword("c"), None),
    "a picture put where a footnote stood": (
        variant(set_pieces("c", text(5) + " ", ("fn",), " " + text(6, 7))), WITH_PICTURE, None),
    "a paragraph with a picture deleted": (variant(PICTURE_PARA), del_para("pic"), None),
    "a paragraph that is only a picture deleted": (variant(PICTURE_ONLY), del_para("pic"), None),
    "a picture taken out of a paragraph": (variant(WITH_PICTURE), set_pieces("c", text(5) + "  " + text(6, 7)), None),
    "a picture swapped for another": (variant(WITH_PICTURE), set_pieces("c", text(5) + " ", ("pic2",), " " + text(6, 7)), None),
    "a table with a picture that lost a row": (
        variant(table("t1", lambda rows: rows[0].__setitem__(0, [[text(30) + " ", ("pic",)]]))), table("t1", del_row), None),
    # fields as a reference manager writes them
    "a word changed next to a citation": (variant(WITH_CITES), reword("c"), None),
    "words changed on both sides of a citation": (
        variant(WITH_CITES), set_pieces("c", text(5)[:-1] + " indeed ", CITE1, " Then " + text(6) + " ", CITE2, " " + text(7)), None),
    "a citation added": (
        variant(WITH_CITES), set_pieces("c", text(5) + " ", CITE1, " " + text(6) + " ", CITE2, " " + text(7) + " ", CITE3), None),
    "a citation deleted": (variant(WITH_CITES), set_pieces("c", text(5) + " ", CITE1, " " + text(6) + " " + text(7)), None),
    "the shown text of a citation changed": (
        variant(WITH_CITES), set_pieces("c", text(5) + " ", CITE1, " " + text(6) + " ", ("cite", 2, "[2, 3]"), " " + text(7)), None),
    "a paragraph with citations deleted": (variant(WITH_CITES), del_para("c"), None),
    "a paragraph with a citation added": (variant(), CITED_PARA, None),
    "a paragraph with citations rewritten": (
        variant(WITH_CITES), set_pieces("c", text(22) + " ", CITE1, " " + text(23) + " ", CITE3), None),
    "a word changed in a caption with a number and a cross-reference": (variant(CAPTION), reword("cap"), None),
    "the result of a cross-reference changed": (variant(CAPTION), both(del_para("cap"), CAPTION2), None),
    "a word changed in a paragraph with a simple field": (variant(WITH_SIMPLE), reword("c"), None),
    "a simple field added to a paragraph": (variant(), WITH_SIMPLE, None),
    "a simple field taken out of a paragraph": (variant(WITH_SIMPLE), set_pieces("c", text(5) + "  " + text(6, 7)), None),
    "a paragraph with a simple field added": (
        variant(), lambda m: m.insert(_at(m, "b"), P("fld", text(20) + " ", ("simple", "PAGE", "3"), ".")), None),
    "a paragraph with a simple field deleted": (
        variant(lambda m: m.insert(_at(m, "b"), P("fld", text(20) + " ", ("simple", "PAGE", "3"), "."))), del_para("fld"), None),
    # a field that runs over several paragraphs
    "an entry added to a bibliography": (variant(BIB), lambda m: m.insert(_at(m, "bib3"), P("bib9", "9. " + text(44))), None),
    "the first entry of a bibliography reworded": (variant(BIB), reword("bib1"), None),
    "the first entry of a bibliography replaced": (variant(BIB), set_pieces("bib1", ("fopen", BIBLIOGRAPHY), "1. " + text(45)), None),
    "the last entry of a bibliography deleted": (
        variant(BIB), both(del_para("bib3"), set_pieces("bib2", "2. " + text(42), ("fend",))), None),
    "a bibliography added": (variant(), BIB, None),
    # signs that are no characters
    "a word changed in a paragraph with a hard hyphen, a symbol and a return": (variant(WITH_SIGNS), reword("c"), None),
    "a symbol taken out": (
        variant(WITH_SIGNS), set_pieces("c", text(5) + " well", ("nbh",), "known " + text(6), ("cr",), text(7), ("shy",)), None),
    "a paragraph with a hard hyphen and a symbol deleted": (variant(WITH_SIGNS), del_para("c"), None),
    # equations: kept where they are the same, refused where they change
    "a word changed in a paragraph with an equation": (variant(WITH_MATH), reword("c"), None),
    "an equation changed": (
        variant(WITH_MATH), set_pieces("c", text(5) + " ", ("math", "x+z"), " " + text(6, 7)), r"\(oMath\) that stands outside a run"),
    "a paragraph with an equation added": (
        variant(), lambda m: m.insert(_at(m, "b"), P("eq", text(20) + " ", ("math", "x+y"))), r"clean paragraph 3 of the body .*\(oMath\)"),
    "a paragraph with an equation deleted": (variant(WITH_MATH), del_para("c"), r"baseline paragraph 5 of the body .*\(oMath\)"),
    # comments
    "a word changed in a paragraph with a comment": (variant(WITH_COMMENT), reword("c"), None),
    "a word changed under a comment": (
        variant(WITH_COMMENT), set_pieces("c", text(5) + " ", ("cm", "a reworded phrase"), " " + text(6, 7)), None),
    "a comment added in a changed paragraph": (variant(), WITH_COMMENT, None),
    "a comment only in the baseline, the paragraph changed": (variant(WITH_COMMENT), set_pieces("c", text(5, 6, 7)), None),
    "a paragraph with a comment deleted": (variant(WITH_COMMENT), del_para("c"), None),
    # a wrapper the writer does not rebuild: the paragraph is shown as deleted and new
    "a word changed in a paragraph with a content control": (variant(WITH_CONTROL), reword("c"), None),
    "a sentence moved to the next paragraph": (variant(), move_sentence("a", "b"), None),
    "a sentence moved across an empty paragraph": (variant(), move_sentence("b", "c"), None),
    "a paragraph rewritten around one sentence": (variant(), rewrite_keeping("a", 1, (20, 21, 22)), None),
    "two paragraphs merged and reworded": (variant(), both(merge_paras("a"), reword("a", 6, "changed")), None),
    "three paragraphs merged": (variant(), both(merge_paras("f"), del_para("h2"), merge_paras("d")), None),
    "a heading deleted and the next paragraph reworded": (variant(), both(del_para("h2"), reword("f")), None),
    "two paragraphs deleted around an empty one": (variant(), both(del_para("b"), del_para("c")), None),
    "an empty paragraph and its neighbour deleted": (variant(), both(del_para("e"), del_para("b")), None),
    "two empty paragraphs added": (variant(), both(add_empty("b"), add_empty("b")), None),
    "an empty paragraph added at the start": (variant(), add_empty("h1"), None),
    "all but one paragraph deleted": (variant(), only("a"), None),
    "everything new": (variant(only("a")), both(replace_para("a", (20, 21)), add_para(None, (22, 23))), None),
    "the only table removed": (variant(), del_para("t1"), None),
    "a table at the start of the document": (variant(to_front("t1")), reword("a"), None),
    "a table added at the start": (variant(), both(add_table("h1"), reword("a")), None),
    "a table at the start removed": (variant(to_front("t1")), del_para("t1"), None),
    "a table at the end of the document": (variant(to_end("t1")), reword("a"), None),
    "a table at the end removed": (variant(to_end("t1")), del_para("t1"), None),
    "a table moved to the end": (variant(), to_end("t1"), None),
    "a table moved between paragraphs": (variant(), move_para("t1", "g"), None),
    "the paragraph between two tables deleted": (variant(add_table("d")), del_para("d"), None),
    "a table and the paragraph before it deleted": (variant(), both(del_para("t1"), del_para("c")), None),
    "a table added and the paragraph before it split": (variant(), both(add_table("d"), split_para("c")), None),
    "a row added and a cell reworded": (variant(), both(table("t1", cell_reword), table("t1", add_row)), None),
    "a cell emptied": (variant(), table("t1", lambda rows: rows[0][0].__setitem__(0, [])), None),
    "an empty cell filled": (variant(table("t1", lambda rows: rows[0][0].__setitem__(0, []))), table("t1", lambda rows: rows[0][0].__setitem__(0, [text(48)])), None),
    "a paragraph of a cell split": (variant(table("t1", lambda rows: rows[0][0].__setitem__(0, [text(30, 35)]))), table("t1", lambda rows: rows[0].__setitem__(0, [[text(30)], [text(35)]])), None),
    "two paragraphs of a cell merged": (variant(), table("t1", lambda rows: rows[1].__setitem__(1, [[text(33, 34)]])), None),
    "a line break added": (variant(), BR("line"), None),
    "a tab removed": (variant(WITH_TAB), set_pieces("c", text(5, 6, 7)), None),
    "a link reworded": (variant(WITH_LINK), set_pieces("c", text(5) + " ", ("link", "a changed phrase"), " " + text(6, 7)), None),
    "a paragraph with a link deleted": (variant(WITH_LINK), del_para("c"), None),
    "a bookmark around text that is reworded": (variant(WITH_MARK), set_pieces("c", text(5) + " ", ("bm", "mark1", "Figure 4"), " " + text(6, 7)), None),
    "a paragraph with a bookmark deleted": (variant(WITH_MARK), del_para("c"), None),
    "a whitespace paragraph deleted": (variant(set_pieces("e", " ")), del_para("e"), None),
    "a paragraph that is only a page break deleted": (variant(set_pieces("e", ("br", "page"))), del_para("e"), None),
    "a paragraph added inside a bookmark over paragraphs": (variant(span("f", "g")), add_para("g", (20, 21)), None),
    "the last paragraph of a bookmark over paragraphs deleted": (variant(span("f", "g")), del_para("g"), None),
    "a paragraph added after a bookmark over paragraphs": (variant(span("f", "g")), add_para(None, (20, 21)), None),
    "paragraphs joined across the end of a bookmark": (variant(span("a", "a")), merge_paras("a"), None),
    "all but one paragraph and the table deleted": (variant(), only("a", "t1"), r"paragraph mark cannot show"),
    "a paragraph added between two tables": (variant(add_table("d")), add_para("t2", (20, 21)), r"paragraph mark cannot show"),
}

# Changes that can be combined freely: each touches another part of the base model.
MIXABLE = [
    reword("a"), reword("f", 5, "revised"), add_para("b", (20, 21)), add_para(None, (24, 25), "tail"),
    del_para("h2"), del_para("g"), split_para("c"), merge_paras("a"), add_empty("d"), del_para("e"),
    table("t1", add_row), table("t1", cell_reword), table("t1", cell_add_para), replace_para("d", (26, 27)),
    add_para("h1", (28,), "lead"),
]
CONFLICTS = [{2, 7}]  # pairs of MIXABLE changes that need the same block


def mixed_cases(count, seed=20311):
    """`count` combinations of two to four changes, the same ones for the same seed."""
    rng = random.Random(seed)
    out = {}
    while len(out) < count:
        picks = sorted(rng.sample(range(len(MIXABLE)), rng.randint(2, 4)))
        if any(pair <= set(picks) for pair in CONFLICTS):
            continue
        # A merge takes the next block, so it runs before the changes that add or remove blocks.
        order = sorted(picks, key=lambda k: (k != 7, k))
        out["mix " + "+".join(map(str, picks))] = (variant(), both(*[MIXABLE[k] for k in order]), None)
    return out


# A baseline with what a manuscript holds besides text, and changes of it that can be combined.
RICH = (WITH_CITES, WITH_NOTE, PICTURE_PARA, CAPTION, BIB)
RICH_MIX = [
    reword("c"), del_para("pic"), reword("cap"), reword("bib1"),
    lambda m: m.insert(_at(m, "bib3"), P("bib9", "9. " + text(44))), del_para("d"), merge_paras("a"),
    set_pieces("c", text(5) + " ", CITE1, " " + text(6) + " ", ("cite", 2, "[2, 3]"), " " + text(7) + " ", CITE3),
    replace_para("c", (26, 27)), table("t1", add_row), add_para("h1", (28,), "lead"), del_para("cap"),
    set_pieces("d", text(8), ("fn2",), " " + text(9) + " ", ("pic2",)),
    lambda m: m.insert(_at(m, "h2"), P("cited", text(20) + " ", CITE3, " " + text(21))), del_para("bib2"),
]


def rich_cases(count, seed=20312):
    """`count` combinations of two to four changes of the rich baseline, the same ones for the same seed."""
    rng = random.Random(seed)
    out = {}
    while len(out) < count:
        picks = sorted(rng.sample(range(len(RICH_MIX)), rng.randint(2, 4)))
        if {2, 11} <= set(picks) or {5, 12} <= set(picks):   # two changes of one block, the first of them a deletion
            continue
        out["rich " + "+".join(map(str, picks))] = (variant(*RICH), both(*[RICH_MIX[k] for k in picks]), None)
    return out


def models(case):
    make, change, refusal = case
    base = make()
    clean = copy.deepcopy(base)
    change(clean)
    return base, clean, refusal
