"""The fake manuscript and supplementary file of the parts acceptance stage
(scripts/e2e/make_parts_fixture.py), and the review and ledger that go with them."""
import importlib.util
import json
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "e2e" / "make_parts_fixture.py"
FIXTURES = ROOT / "tests" / "fixtures"


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


fx = load("make_parts_fixture", SCRIPT)

try:
    import docx  # noqa: F401

    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False


class CommittedFixtures(unittest.TestCase):
    """The committed files, read with the standard library only."""

    @classmethod
    def setUpClass(cls):
        cls.paper = fx.dump(FIXTURES / "parts_paper.docx")
        cls.si = fx.dump(FIXTURES / "parts_si.docx")

    def test_the_texts_are_those_of_the_script(self):
        self.assertEqual([t for t in self.paper if t], [text for _, text in fx.PAPER])
        cells = [cell for row in fx.SI_TABLE for cell in row]
        self.assertEqual([t for t in self.si if t], [text for _, text in fx.SI] + cells)

    def test_the_forbidden_word_is_in_neither_document(self):
        word = re.compile(r"\bnovel(?:ty)?\b", re.I)
        for name, texts in (("paper", self.paper), ("si", self.si)):
            self.assertEqual([t for t in texts if word.search(t)], [], name)
        # ... and the pattern of the ledger is that word.
        ledger = json.loads((FIXTURES / "parts_ledger.json").read_text(encoding="utf-8"))
        self.assertEqual(len(ledger["items"]), 1)
        item = ledger["items"][0]
        self.assertEqual((item["kind"], item["check"]["type"]), ("forbid", "regex"))
        rule = re.compile(item["check"]["pattern"], re.I)
        self.assertTrue(rule.search(fx.GUARD["replace"]))
        self.assertFalse(rule.search(fx.GUARD["find"]))
        self.assertEqual(fx.SPEC["forbidden_word"], "novel")

    def test_the_guard_edit_has_one_place_in_the_supplementary_file(self):
        whole = "\n".join(self.si)
        self.assertEqual(whole.count(fx.GUARD["find"]), 1)
        self.assertEqual(fx.GUARD["ref"], fx.SPEC["refs"]["si_only"])

    def test_each_comment_is_about_a_sentence_of_the_file_it_names(self):
        paper = "\n".join(self.paper)
        si = "\n".join(self.si)
        # R1.1: the supplementary methods lack the temperature; the table has it.
        self.assertEqual(si.count("Starter cultures were incubated overnight"), 1)
        self.assertNotIn("incubated", paper)
        self.assertIn("30 degrees Celsius", si)
        # R1.2: the shaking speed is missing in both files, once each.
        self.assertEqual(paper.count("at a constant speed"), 1)
        self.assertEqual(si.count("at a constant speed"), 1)
        self.assertNotIn("rpm", paper + si)
        # R1.3: only the manuscript has the slip.
        self.assertEqual(paper.count("grew good"), 1)
        self.assertNotIn("grew good", si)
        # R1.4: the figure is named in the manuscript and is no part of either file.
        self.assertIn("Figure 1", paper)
        self.assertNotIn("Figure 1", si)

    def test_the_review_has_four_comments_and_gives_the_numbers_it_asks_for(self):
        review = (FIXTURES / "parts_review.txt").read_text(encoding="utf-8")
        self.assertEqual(re.findall(r"^Comment (\d+):", review, re.M), ["1", "2", "3", "4"])
        self.assertRegex(review, r"(?m)^Reviewer #1: Comments to the Author$")
        # A number an edit writes must be in the comment (the tool refuses others).
        self.assertIn("30 degrees Celsius", review)
        self.assertIn("180 rpm", review)
        self.assertNotRegex(review, r"(?i)\bnovel")
        self.assertNotIn("—", review)
        # The spec names the four comments and gives requests for each.
        self.assertEqual(sorted(fx.SPEC["refs"].values()), ["R1.1", "R1.2", "R1.3", "R1.4"])
        self.assertEqual(sorted(fx.SPEC["requirements"]), ["R1.1", "R1.2", "R1.3", "R1.4"])
        self.assertEqual(len(fx.SPEC["requirements"][fx.SPEC["refs"]["both"]]), 2)

    def test_spec_and_dump_from_the_command_line(self):
        out = subprocess.run(
            [sys.executable, str(SCRIPT), "--spec"], capture_output=True, text=True, check=True
        )
        self.assertEqual(json.loads(out.stdout), fx.SPEC)
        out = subprocess.run(
            [sys.executable, str(SCRIPT), "--dump", str(FIXTURES / "parts_si.docx")],
            capture_output=True,
            text=True,
            check=True,
        )
        self.assertEqual(json.loads(out.stdout), self.si)


@unittest.skipUnless(HAVE_DOCX, "python-docx is not installed")
class BuiltFixtures(unittest.TestCase):
    def test_a_new_build_has_the_text_of_the_committed_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            for name in ("paper", "si"):
                built = fx.build(name, Path(tmp) / f"{name}.docx")
                self.assertEqual(fx.dump(built), fx.dump(fx.OUT[name]), name)

    def test_the_supplementary_file_has_one_table(self):
        import docx as _docx

        with tempfile.TemporaryDirectory() as tmp:
            doc = _docx.Document(str(fx.build("si", Path(tmp) / "si.docx")))
            self.assertEqual(len(doc.tables), 1)
            self.assertEqual(
                [[c.text for c in row.cells] for row in doc.tables[0].rows],
                [list(r) for r in fx.SI_TABLE],
            )


if __name__ == "__main__":
    unittest.main()
